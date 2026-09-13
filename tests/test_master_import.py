import io
from datetime import datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

from app import cost_increase, master_import, predicted_increase, storage
from app import passport as passport_module


def _png_bytes(color="red"):
    buf = io.BytesIO()
    PILImage.new("RGB", (4, 4), color).save(buf, format="PNG")
    return buf.getvalue()


def _portfolio_workbook(shared_rows, blocks):
    """A workbook shaped like the real portfolio file.

    ``shared_rows`` is ``[(row, number, label), ...]`` for the "№" / work
    type columns — written once, since every project block reads the same
    physical cells for its row labels (only the cost columns differ).

    Each entry in ``blocks`` is ``(name, passport_row_values, cost_headers,
    cost_values)``:
    - ``passport_row_values`` maps row number -> value (year/class/GC/areas)
    - ``cost_headers`` is the list of row-24 header strings for this block
    - ``cost_values`` maps row -> list of values, aligned 1:1 with
      ``cost_headers``
    """
    wb = Workbook()
    ws = wb.active
    ws.title = master_import.SHEET_NAME

    for row, number, label in shared_rows:
        ws.cell(row=row, column=master_import.NUMBER_COLUMN, value=number)
        ws.cell(row=row, column=master_import.LABEL_COLUMN, value=label)

    # Real project blocks start two columns past the shared label column
    # (column 4 sits empty in between) — not at FIRST_PROJECT_COLUMN itself,
    # which only marks where the *search* for block starts begins.
    col = master_import.LABEL_COLUMN + 2
    for name, passport_rows, cost_headers, cost_values in blocks:
        ws.cell(row=master_import.PROJECT_NAME_ROW, column=col, value=name)
        for row, value in passport_rows.items():
            ws.cell(row=row, column=col, value=value)
        for offset, header in enumerate(cost_headers):
            ws.cell(row=master_import.ROW_VERSION_HEADER, column=col + offset, value=header)
        for row, values in cost_values.items():
            for offset, value in enumerate(values):
                ws.cell(row=row, column=col + offset, value=value)
        col += len(cost_headers) or 10
    return wb


def _save_and_reload(wb, tmp_path, name="master.xlsx"):
    path = tmp_path / name
    wb.save(path)
    return path


def test_parses_two_project_blocks(tmp_path):
    row = master_import.FIRST_COST_ROW
    shared_rows = [(row, "1", 'Разработка стадии "Р"')]
    blocks = [
        (
            "MIRA\nПроспект Мира",
            {
                master_import.ROW_YEAR_SIGNED: datetime(2025, 2, 20),
                master_import.ROW_BUILDING_CLASS: "Жилая недвижимость (Бизнес)",
                master_import.ROW_GENERAL_CONTRACTOR: "АНТТЕК",
                master_import.ROW_UNDERGROUND_AREA: 13341.3,
                master_import.ROW_ABOVEGROUND_AREA: 62399.7,
                master_import.ROW_TOTAL_AREA: 75740.9,
            },
            ["Объем", "Протокол ОУ", "Стоимость\nна 1 м² ЖК", "ДГП"],
            {row: [1000, 500, 50.5, 600]},
        ),
        (
            "Nicole 1",
            {
                master_import.ROW_YEAR_SIGNED: " Март 2024",
                master_import.ROW_BUILDING_CLASS: "Жилая недвижимость (Делюкс)",
                master_import.ROW_GENERAL_CONTRACTOR: "АО ФОДД",
                master_import.ROW_UNDERGROUND_AREA: 100.0,
                master_import.ROW_ABOVEGROUND_AREA: 200.0,
                master_import.ROW_TOTAL_AREA: 300.0,
            },
            ["Объем", "Протокол ОУ"],
            {row: [10, 20]},
        ),
    ]
    wb = _portfolio_workbook(shared_rows, blocks)
    path = _save_and_reload(wb, tmp_path)

    projects = master_import.parse_workbook(path)

    assert len(projects) == 2
    mira, nicole = projects

    assert mira["raw_name"] == "MIRA / Проспект Мира"
    assert mira["name_candidates"] == ["MIRA", "Проспект Мира"]
    assert mira["passport"] == {
        "year_signed": "2025",
        "building_class": "Бизнес",
        "general_contractor": "АНТТЕК",
        "underground_area_sqm": 13341.3,
        "aboveground_area_sqm": 62399.7,
        "total_area_sqm": 75740.9,
        # No "Итого"/"2"/"2.1"/"2.2" rows in this fixture's cost table, so
        # nothing to fill in for any of these.
        "contract_price_rub": None,
        "smr_term": None,
        "advance_payment": None,
        "performance_bond_pct": None,
        "bank_guarantee": None,
        "vat": None,
    }
    assert mira["cost_table"]["primary_version"] == "Протокол ОУ"
    assert mira["cost_table"]["columns"] == [
        "Объем", "Протокол ОУ", "Протокол ОУ, руб/м²", "ДГП",
    ]
    assert mira["cost_table"]["rows"] == [{
        "number": "1",
        "label": 'Разработка стадии "Р"',
        "values": {
            "Объем": 1000, "Протокол ОУ": 500, "Протокол ОУ, руб/м²": 50.5, "ДГП": 600,
        },
    }]

    assert nicole["raw_name"] == "Nicole 1"
    assert nicole["name_candidates"] == ["Nicole 1"]
    assert nicole["passport"]["year_signed"] == "2024"
    assert nicole["passport"]["building_class"] == "Делюкс"
    assert nicole["cost_table"]["columns"] == ["Объем", "Протокол ОУ"]
    assert nicole["cost_table"]["rows"] == [{
        "number": "1",
        "label": 'Разработка стадии "Р"',
        "values": {"Объем": 10, "Протокол ОУ": 20},
    }]


def test_contract_price_comes_from_the_итого_row_not_a_sum_of_line_items(tmp_path):
    # Real case (БЦ JOIS / Силикатный): summing every numbered line gives
    # 18 553 243 164.22, but the sheet's own «Итого» row says
    # 18 464 998 172.15 — some line items (bank guarantees here) aren't
    # folded into the printed total, so only that row's own value is right.
    row1 = master_import.FIRST_COST_ROW
    row_guarantee = row1 + 1
    row_total = row1 + 2
    shared_rows = [
        (row1, "1", "Разработка стадии \"Р\""),
        (row_guarantee, "2.2", "Банковская гарантия на исполнение контракта"),
        (row_total, None, "Итого СМР, руб. с НДС 20%"),
    ]
    blocks = [(
        "БЦ JOIS\nСиликатный",
        {},
        ["Протокол ОУ", "ДГП"],
        {
            row1: [18_000_000, 17_000_000],
            row_guarantee: [500_000, None],
            row_total: [18_000_000, 17_000_000],
        },
    )]
    wb = _portfolio_workbook(shared_rows, blocks)
    path = _save_and_reload(wb, tmp_path)

    projects = master_import.parse_workbook(path)

    assert projects[0]["passport"]["contract_price_rub"] == 18_000_000


def test_contract_price_takes_the_last_of_several_итого_rows(tmp_path):
    row1 = master_import.FIRST_COST_ROW
    row_subtotal = row1 + 1
    row_extra = row1 + 2
    row_final_total = row1 + 3
    shared_rows = [
        (row1, "1", "Разработка стадии \"Р\""),
        (row_subtotal, None, "Итого СМР, руб. с НДС 20%"),
        (row_extra, "16", "SHELL & CORE"),
        (row_final_total, None, "Итого СМР, в тч Отделка и Нулевой цикл, руб. с НДС с 20%"),
    ]
    blocks = [(
        "MIRA",
        {},
        ["Протокол ОУ"],
        {
            row1: [1_000_000],
            row_subtotal: [1_000_000],
            row_extra: [200_000],
            row_final_total: [1_200_000],
        },
    )]
    wb = _portfolio_workbook(shared_rows, blocks)
    path = _save_and_reload(wb, tmp_path)

    projects = master_import.parse_workbook(path)

    assert projects[0]["passport"]["contract_price_rub"] == 1_200_000


def test_contract_price_is_none_without_an_итого_row(tmp_path):
    row = master_import.FIRST_COST_ROW
    shared_rows = [(row, "1", "Разработка стадии \"Р\"")]
    blocks = [("MIRA", {}, ["Протокол ОУ"], {row: [500]})]
    wb = _portfolio_workbook(shared_rows, blocks)
    path = _save_and_reload(wb, tmp_path)

    projects = master_import.parse_workbook(path)

    assert projects[0]["passport"]["contract_price_rub"] is None


def _block_with_items_2(row1_extra=None, row21_volume=None, row22_volume=None, row2_volume="29 мес в тч ЗОС"):
    row1 = master_import.FIRST_COST_ROW
    row2 = row1 + 1
    row21 = row1 + 2
    row22 = row1 + 3
    shared_rows = [
        (row1, "1", "Разработка стадии \"Р\""),
        (row2, "2", "Подготовительные работы и содержание площадки"),
        (row21, "2.1", "Банковская гарантия на авансовые платежи"),
        (row22, "2.2", "Банковкская гаранития на исполнение контракта"),
    ]
    cost_values = {
        row1: [500],
        row2: [row2_volume],
        row21: [row21_volume],
        row22: [row22_volume],
    }
    if row1_extra is not None:
        cost_values[row1] = row1_extra
    blocks = [("MIRA", {}, ["Объем"], cost_values)]
    return shared_rows, blocks


def test_smr_term_and_percentages_come_from_item_2_and_its_sub_items(tmp_path):
    shared_rows, blocks = _block_with_items_2(row21_volume=0.03, row22_volume=0.025)
    wb = _portfolio_workbook(shared_rows, blocks)
    path = _save_and_reload(wb, tmp_path)

    passport = master_import.parse_workbook(path)[0]["passport"]

    assert passport["smr_term"] == "29"
    assert passport["advance_payment"] == "3%"
    assert passport["performance_bond_pct"] == "2.5%"


def test_advance_payment_is_written_as_is_when_the_cell_is_текстовый_статус(tmp_path):
    # Real case: "да"/"нет"/"предоставлено" in the advance-payment
    # guarantee row is itself a meaningful answer, so it's written down
    # exactly as the sheet has it — not converted to a percent, not dropped.
    shared_rows, blocks = _block_with_items_2(row21_volume="нет")
    wb = _portfolio_workbook(shared_rows, blocks)
    path = _save_and_reload(wb, tmp_path)

    passport = master_import.parse_workbook(path)[0]["passport"]

    assert passport["advance_payment"] == "нет"


def test_performance_bond_is_none_when_the_cell_is_текстовый_статус(tmp_path):
    # Unlike advance_payment, this field is specifically a percentage — a
    # "да"/"нет" placeholder means the sheet has no figure there yet.
    shared_rows, blocks = _block_with_items_2(row22_volume="да")
    wb = _portfolio_workbook(shared_rows, blocks)
    path = _save_and_reload(wb, tmp_path)

    passport = master_import.parse_workbook(path)[0]["passport"]

    assert passport["performance_bond_pct"] is None


@pytest.mark.parametrize("volume, expected", [
    (0.03, "Включено"),
    ("нет", "Не включено"),
    ("да", "Включено"),
    ("предоставлено", "предоставлено"),
    (None, None),
])
def test_bank_guarantee_reads_the_same_cell_as_advance_payment(tmp_path, volume, expected):
    shared_rows, blocks = _block_with_items_2(row21_volume=volume)
    wb = _portfolio_workbook(shared_rows, blocks)
    path = _save_and_reload(wb, tmp_path)

    passport = master_import.parse_workbook(path)[0]["passport"]

    assert passport["bank_guarantee"] == expected


def test_vat_is_read_from_the_итого_row_label(tmp_path):
    row1 = master_import.FIRST_COST_ROW
    row_total = row1 + 1
    shared_rows = [
        (row1, "1", "Разработка стадии \"Р\""),
        (row_total, None, "Итого СМР, руб. с НДС 20%"),
    ]
    blocks = [("MIRA", {}, ["Протокол ОУ"], {row1: [500], row_total: [500]})]
    wb = _portfolio_workbook(shared_rows, blocks)
    path = _save_and_reload(wb, tmp_path)

    passport = master_import.parse_workbook(path)[0]["passport"]

    assert passport["vat"] == "20%"


def test_vat_is_none_without_an_итого_row(tmp_path):
    row = master_import.FIRST_COST_ROW
    shared_rows = [(row, "1", "Разработка стадии \"Р\"")]
    blocks = [("MIRA", {}, ["Протокол ОУ"], {row: [500]})]
    wb = _portfolio_workbook(shared_rows, blocks)
    path = _save_and_reload(wb, tmp_path)

    passport = master_import.parse_workbook(path)[0]["passport"]

    assert passport["vat"] is None


def test_smr_term_rejects_an_implausible_month_count(tmp_path):
    # Real case (Jois): item 2's «Объем» cell held a stray cost figure
    # (billions) instead of a month count.
    shared_rows, blocks = _block_with_items_2(row2_volume=2089963755.27)
    wb = _portfolio_workbook(shared_rows, blocks)
    path = _save_and_reload(wb, tmp_path)

    passport = master_import.parse_workbook(path)[0]["passport"]

    assert passport["smr_term"] is None


def test_contract_terms_fields_never_overwrite_an_existing_value(tmp_path):
    slug = storage.create_project(tmp_path, "MIRA")
    existing = passport_module.build_passport("MIRA")
    existing["smr_term"] = "33"
    passport_module.save_passport(existing, storage.passport_path(tmp_path, slug))

    parsed = [{
        "raw_name": "MIRA",
        "name_candidates": ["MIRA"],
        "passport": {
            "year_signed": None, "building_class": None, "general_contractor": None,
            "underground_area_sqm": None, "aboveground_area_sqm": None, "total_area_sqm": None,
            "contract_price_rub": None, "smr_term": "29", "advance_payment": "3%",
            "performance_bond_pct": None,
        },
        "cost_table": {"primary_version": "Протокол ОУ", "columns": [], "rows": []},
    }]

    master_import.apply_import(tmp_path, parsed)

    data = passport_module.load_passport(storage.passport_path(tmp_path, slug))
    assert data["smr_term"] == "33"
    assert data["advance_payment"] == "3%"


def test_parses_cover_photo_anchored_at_project_start(tmp_path):
    shared_rows = [(master_import.FIRST_COST_ROW, "1", "Фундамент")]
    blocks = [
        (
            "MIRA\nПроспект Мира",
            {master_import.ROW_YEAR_SIGNED: datetime(2025, 2, 20)},
            ["Протокол ОУ"],
            {master_import.FIRST_COST_ROW: [500]},
        ),
        (
            "Nicole 1",
            {master_import.ROW_YEAR_SIGNED: datetime(2024, 6, 1)},
            ["Протокол ОУ"],
            {master_import.FIRST_COST_ROW: [20]},
        ),
    ]
    wb = _portfolio_workbook(shared_rows, blocks)
    ws = wb[master_import.SHEET_NAME]
    first_col = master_import.LABEL_COLUMN + 2  # MIRA's own start column, matching _portfolio_workbook
    photo = XLImage(io.BytesIO(_png_bytes()))
    # A cell reference at column ``first_col`` (1-indexed) anchors at
    # 0-indexed column ``first_col - 1`` — the "one column before its own
    # start" case observed in the real file.
    ws.add_image(photo, f"{get_column_letter(first_col)}{master_import.PROJECT_NAME_ROW + 1}")
    path = _save_and_reload(wb, tmp_path)

    projects = master_import.parse_workbook(path)

    mira, nicole = projects
    assert mira["cover"] is not None
    assert mira["cover"]["ext"] == ".png"
    assert mira["cover"]["data"].startswith(b"\x89PNG")
    assert nicole["cover"] is None


def test_missing_sheet_raises(tmp_path):
    wb = Workbook()
    path = _save_and_reload(wb, tmp_path)
    with pytest.raises(master_import.MasterImportError):
        master_import.parse_workbook(path)


def test_not_an_xlsx_raises(tmp_path):
    path = tmp_path / "not_excel.xlsx"
    path.write_bytes(b"this is not a zip file")
    with pytest.raises(master_import.MasterImportError):
        master_import.parse_workbook(path)


@pytest.mark.parametrize("text, expected", [
    ("Жилая недвижимость (Бизнес)", "Бизнес"),
    ("Жилая недвижимость (Бизнес-премиум)", "Бизнес - Премиум"),
    ("Жилая недвижимость (Комфорт)", "Комфорт"),
    ("Офисная недвижимость (класс А)", None),
    ("Складская недвижимость", None),
    (None, None),
])
def test_normalize_building_class(text, expected):
    assert master_import.normalize_building_class(text) == expected


@pytest.mark.parametrize("value, expected", [
    (datetime(2025, 2, 20), "2025"),
    (" Март 2024", "2024"),
    (44440, None),  # out of the plausible-year range, not a real year
    (2021, "2021"),
    (None, None),
    ("no year here", None),
])
def test_extract_year(value, expected):
    assert master_import.extract_year(value) == expected


@pytest.mark.parametrize("candidates, existing, expected_matches", [
    (["MIRA", "Проспект Мира"], "Проспект Мира", True),
    (["ТУШИНО 1", "ЖК CITYZEN"], "CITYZEN", True),
    (["ЖК JOIS 1", "Силикатный"], "JOIS 1", True),
    (["Nicole 1"], "Nicole 2", False),
    (["Слава 2,3"], "Совсем другой проект", False),
])
def test_match_existing_slug(candidates, existing, expected_matches):
    project_names = {"slug-1": existing}
    result = master_import.match_existing_slug(candidates, project_names)
    assert result == ("slug-1" if expected_matches else None)


def test_match_existing_slug_no_projects():
    assert master_import.match_existing_slug(["Anything"], {}) is None


def _passport_with_name(name):
    data = passport_module.build_passport(name)
    return data


def test_apply_import_creates_new_project(tmp_path):
    parsed = [{
        "raw_name": "MIRA / Проспект Мира",
        "name_candidates": ["MIRA", "Проспект Мира"],
        "passport": {
            "year_signed": "2025", "building_class": "Бизнес",
            "general_contractor": "АНТТЕК", "underground_area_sqm": 100.0,
            "aboveground_area_sqm": 200.0, "total_area_sqm": 300.0,
        },
        "cost_table": {"primary_version": "Протокол ОУ", "columns": [], "rows": []},
    }]

    report = master_import.apply_import(tmp_path, parsed)

    assert len(report) == 1
    assert report[0]["action"] == "created"
    slug = report[0]["slug"]
    data = passport_module.load_passport(storage.passport_path(tmp_path, slug))
    assert data["project_name"] == "MIRA / Проспект Мира"
    assert data["year_signed"] == "2025"
    assert data["building_class"] == "Бизнес"
    assert data["general_contractor"] == "АНТТЕК"
    assert data["total_area_sqm"] == 300.0
    assert data["address"] is None
    assert storage.master_import_path(tmp_path, slug).exists()


def test_apply_import_saves_cover_photo_when_present(tmp_path):
    parsed = [{
        "raw_name": "MIRA / Проспект Мира",
        "name_candidates": ["MIRA", "Проспект Мира"],
        "passport": {
            "year_signed": None, "building_class": None, "general_contractor": None,
            "underground_area_sqm": None, "aboveground_area_sqm": None, "total_area_sqm": None,
        },
        "cost_table": {"primary_version": "Протокол ОУ", "columns": [], "rows": []},
        "cover": {"data": _png_bytes(), "ext": ".png"},
    }]

    report = master_import.apply_import(tmp_path, parsed)

    cover_path = storage.cover_path(tmp_path, report[0]["slug"])
    assert cover_path is not None
    assert cover_path.suffix == ".png"
    assert cover_path.read_bytes().startswith(b"\x89PNG")


def test_apply_import_without_cover_key_does_not_fail(tmp_path):
    # Hand-built parsed entries (as opposed to ones from parse_workbook)
    # may omit "cover" entirely — apply_import must not require it.
    parsed = [{
        "raw_name": "Без фото",
        "name_candidates": ["Без фото"],
        "passport": {
            "year_signed": None, "building_class": None, "general_contractor": None,
            "underground_area_sqm": None, "aboveground_area_sqm": None, "total_area_sqm": None,
        },
        "cost_table": {"primary_version": "Протокол ОУ", "columns": [], "rows": []},
    }]

    report = master_import.apply_import(tmp_path, parsed)

    assert storage.cover_path(tmp_path, report[0]["slug"]) is None


def test_apply_import_updates_existing_project_without_touching_address(tmp_path):
    slug = storage.create_project(tmp_path, "Проспект Мира")
    existing = _passport_with_name("Проспект Мира")
    existing["address"] = "г. Москва, вручную вписанный адрес"
    existing["year_signed"] = "2020"  # to be overwritten
    passport_module.save_passport(existing, storage.passport_path(tmp_path, slug))

    parsed = [{
        "raw_name": "MIRA / Проспект Мира",
        "name_candidates": ["MIRA", "Проспект Мира"],
        "passport": {
            "year_signed": "2025", "building_class": "Бизнес",
            "general_contractor": "АНТТЕК", "underground_area_sqm": 100.0,
            "aboveground_area_sqm": 200.0, "total_area_sqm": 300.0,
        },
        "cost_table": {"primary_version": "Протокол ОУ", "columns": [], "rows": []},
    }]

    report = master_import.apply_import(tmp_path, parsed)

    assert report[0]["action"] == "updated"
    assert report[0]["slug"] == slug
    data = passport_module.load_passport(storage.passport_path(tmp_path, slug))
    assert data["year_signed"] == "2025"
    assert data["address"] == "г. Москва, вручную вписанный адрес"
    assert data["project_name"] == "Проспект Мира"


def test_apply_import_does_not_merge_two_new_projects_sharing_a_name_line(tmp_path):
    # Real case from the portfolio file: "ВЕРЕЙСКАЯ UB2 / ЖК SET" and
    # "ВЕРЕЙСКАЯ UB3c / ЖК SET" are two different buildings that share their
    # second name line. On an empty root, both must become new projects —
    # matching the second against the first (freshly created a moment
    # earlier, in the same run) would silently lose one of them.
    parsed = [
        {
            "raw_name": "ВЕРЕЙСКАЯ UB2 / ЖК SET",
            "name_candidates": ["ВЕРЕЙСКАЯ UB2", "ЖК SET"],
            "passport": {
                "year_signed": "2024", "building_class": None, "general_contractor": None,
                "underground_area_sqm": None, "aboveground_area_sqm": None,
                "total_area_sqm": None,
            },
            "cost_table": {"primary_version": "Протокол ОУ", "columns": [], "rows": []},
        },
        {
            "raw_name": "ВЕРЕЙСКАЯ UB3c / ЖК SET",
            "name_candidates": ["ВЕРЕЙСКАЯ UB3c", "ЖК SET"],
            "passport": {
                "year_signed": "2025", "building_class": None, "general_contractor": None,
                "underground_area_sqm": None, "aboveground_area_sqm": None,
                "total_area_sqm": None,
            },
            "cost_table": {"primary_version": "Протокол ОУ", "columns": [], "rows": []},
        },
    ]

    report = master_import.apply_import(tmp_path, parsed)

    assert [r["action"] for r in report] == ["created", "created"]
    assert len(set(r["slug"] for r in report)) == 2
    assert len(storage.list_project_slugs(tmp_path)) == 2


def test_apply_import_leaves_none_fields_untouched(tmp_path):
    slug = storage.create_project(tmp_path, "Проспект Мира")
    existing = _passport_with_name("Проспект Мира")
    existing["general_contractor"] = "ООО «АНТТЕК»"
    passport_module.save_passport(existing, storage.passport_path(tmp_path, slug))

    parsed = [{
        "raw_name": "Проспект Мира",
        "name_candidates": ["Проспект Мира"],
        "passport": {
            "year_signed": None, "building_class": None,
            "general_contractor": None, "underground_area_sqm": None,
            "aboveground_area_sqm": None, "total_area_sqm": None,
        },
        "cost_table": {"primary_version": "Протокол ОУ", "columns": [], "rows": []},
    }]

    master_import.apply_import(tmp_path, parsed)

    data = passport_module.load_passport(storage.passport_path(tmp_path, slug))
    assert data["general_contractor"] == "ООО «АНТТЕК»"


# --- Reading the saved cost table back for смета/удорожание/coefficients ---

def _cost_table(rows):
    return {
        "primary_version": "Протокол ОУ",
        "columns": ["Объем", "Протокол ОУ", "ДГП", "ИТОГО ДС", "Предполагаемое ДС"],
        "rows": rows,
    }


_SAMPLE_ROWS = [
    {
        "number": "5", "label": "Ж/Б конструкции",
        "values": {"Объем": 65536.05, "ДГП": 3922867565.29, "ИТОГО ДС": 4000000000.0,
                   "Предполагаемое ДС": 4100000000.0},
    },
    {
        "number": "6", "label": "Металлические конструкции",
        "values": {"Объем": 161.6, "ДГП": 50907060.14, "ИТОГО ДС": 51000000.0,
                   "Предполагаемое ДС": None},
    },
    {
        "number": "8", "label": "Фасад",
        "values": {"Объем": 94342.98, "ДГП": 4630308322.51, "ИТОГО ДС": 4700000000.0,
                   "Предполагаемое ДС": 4800000000.0},
    },
    {
        "number": "2.1", "label": "Банковская гарантия на авансовые платежи",
        "values": {"Объем": "нет", "ДГП": None, "ИТОГО ДС": None, "Предполагаемое ДС": None},
    },
    {
        "number": None, "label": "Итого СМР, руб. с НДС с 20%",
        "values": {"Объем": None, "ДГП": 8604082947.94, "ИТОГО ДС": 8751000000.0,
                   "Предполагаемое ДС": 8900000000.0},
    },
]


def test_estimate_sections_from_cost_table_reads_the_дгп_column(tmp_path):
    sections = master_import.estimate_sections_from_cost_table(_cost_table(_SAMPLE_ROWS))

    by_key = {s.key: s.amount for s in sections}
    # Ж/Б конструкции and Металлические конструкции both classify as
    # "concrete" — kept as two Section entries here (routes.py/excel_report
    # sum them), matching how a real смета with two such rows would look.
    assert sorted(s.key for s in sections) == ["concrete", "concrete", "facade"]
    assert by_key["facade"] == Decimal("4630308322.51")
    # The bank-guarantee row and the «Итого» row both classify() to None,
    # same as they would coming from a real смета — left out entirely.
    assert not any("Банковская" in s.name for s in sections)
    assert not any("Итого" in s.name for s in sections)


def test_concrete_volume_from_cost_table_is_the_жб_rows_own_объем(tmp_path):
    volume = master_import.concrete_volume_from_cost_table(_cost_table(_SAMPLE_ROWS))
    # Not the "Металлические конструкции" row, which shares the "concrete"
    # category but has no concrete volume of its own (161.6 would be
    # implausibly small next to 65536.05 m³, a different unit entirely).
    assert volume == 65536.05


def test_facade_area_from_cost_table_is_the_фасад_rows_own_объем(tmp_path):
    area = master_import.facade_area_from_cost_table(_cost_table(_SAMPLE_ROWS))
    assert area == 94342.98


def test_cost_increase_lines_from_cost_table_reads_итого_дс(tmp_path):
    lines = master_import.cost_increase_lines_from_cost_table(_cost_table(_SAMPLE_ROWS))

    assert all(isinstance(line, cost_increase.AmountLine) for line in lines)
    amounts = {line.name: line.amount for line in lines}
    assert amounts["Ж/Б конструкции"] == Decimal("4000000000.0")
    assert amounts["Металлические конструкции"] == Decimal("51000000.0")
    assert "Банковская гарантия на авансовые платежи" not in amounts


def test_predicted_increase_lines_from_cost_table_reads_предполагаемое_дс(tmp_path):
    lines = master_import.predicted_increase_lines_from_cost_table(_cost_table(_SAMPLE_ROWS))

    assert all(isinstance(line, predicted_increase.Line) for line in lines)
    amounts = {line.name: line.amount for line in lines}
    assert amounts["Ж/Б конструкции"] == Decimal("4100000000.0")
    # None in the source cell (no предполагаемое ДС for this row yet) means
    # no line for it — not a line worth zero.
    assert "Металлические конструкции" not in amounts


def test_load_cost_table_round_trips_what_apply_import_wrote(tmp_path):
    parsed = [{
        "raw_name": "MIRA", "name_candidates": ["MIRA"],
        "passport": {
            "year_signed": None, "building_class": None, "general_contractor": None,
            "underground_area_sqm": None, "aboveground_area_sqm": None, "total_area_sqm": None,
            "contract_price_rub": None, "smr_term": None, "advance_payment": None,
            "performance_bond_pct": None, "bank_guarantee": None, "vat": None,
        },
        "cost_table": _cost_table(_SAMPLE_ROWS),
    }]
    report = master_import.apply_import(tmp_path, parsed)

    loaded = master_import.load_cost_table(tmp_path, report[0]["slug"])
    assert loaded == _cost_table(_SAMPLE_ROWS)


def test_load_cost_table_is_none_when_never_imported(tmp_path):
    slug = storage.create_project(tmp_path, "Без импорта")
    assert master_import.load_cost_table(tmp_path, slug) is None
