import io

from openpyxl import Workbook

from app import create_app, master_import, passport as passport_module, storage


def _workbook_bytes(name, year, building_class, gc, areas):
    wb = Workbook()
    ws = wb.active
    ws.title = master_import.SHEET_NAME
    col = 5
    ws.cell(row=master_import.PROJECT_NAME_ROW, column=col, value=name)
    ws.cell(row=master_import.ROW_YEAR_SIGNED, column=col, value=year)
    ws.cell(row=master_import.ROW_BUILDING_CLASS, column=col, value=building_class)
    ws.cell(row=master_import.ROW_GENERAL_CONTRACTOR, column=col, value=gc)
    under, above, total = areas
    ws.cell(row=master_import.ROW_UNDERGROUND_AREA, column=col, value=under)
    ws.cell(row=master_import.ROW_ABOVEGROUND_AREA, column=col, value=above)
    ws.cell(row=master_import.ROW_TOTAL_AREA, column=col, value=total)
    ws.cell(row=master_import.ROW_VERSION_HEADER, column=col, value="Протокол ОУ")
    ws.cell(row=master_import.FIRST_COST_ROW, column=master_import.NUMBER_COLUMN, value="1")
    ws.cell(row=master_import.FIRST_COST_ROW, column=master_import.LABEL_COLUMN, value="Фундамент")
    ws.cell(row=master_import.FIRST_COST_ROW, column=col, value=1000)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_master_import_form_loads(tmp_path):
    app = create_app(tmp_path)
    client = app.test_client()
    resp = client.get("/master-import")
    assert resp.status_code == 200


def test_run_master_import_creates_project(tmp_path):
    app = create_app(tmp_path)
    client = app.test_client()
    data = _workbook_bytes(
        "MIRA\nПроспект Мира", None, "Жилая недвижимость (Бизнес)", "АНТТЕК",
        (100.0, 200.0, 300.0),
    )

    resp = client.post(
        "/master-import",
        data={"workbook_file": (io.BytesIO(data), "portfolio.xlsx")},
        content_type="multipart/form-data",
    )
    body = resp.get_data(as_text=True)

    assert resp.status_code == 200
    assert "Создан" in body
    [slug] = storage.list_project_slugs(tmp_path)
    passport_data = passport_module.load_passport(storage.passport_path(tmp_path, slug))
    assert passport_data["building_class"] == "Бизнес"
    assert passport_data["general_contractor"] == "АНТТЕК"
    assert passport_data["total_area_sqm"] == 300.0
    assert storage.master_import_path(tmp_path, slug).exists()


def test_run_master_import_updates_existing_project(tmp_path):
    app = create_app(tmp_path)
    client = app.test_client()
    slug = storage.create_project(tmp_path, "Проспект Мира")
    existing = passport_module.build_passport("Проспект Мира")
    existing["address"] = "вручную вписанный адрес"
    passport_module.save_passport(existing, storage.passport_path(tmp_path, slug))

    data = _workbook_bytes(
        "MIRA\nПроспект Мира", None, "Жилая недвижимость (Бизнес)", "АНТТЕК",
        (100.0, 200.0, 300.0),
    )
    resp = client.post(
        "/master-import",
        data={"workbook_file": (io.BytesIO(data), "portfolio.xlsx")},
        content_type="multipart/form-data",
    )
    body = resp.get_data(as_text=True)

    assert resp.status_code == 200
    assert "Обновлён" in body
    assert len(storage.list_project_slugs(tmp_path)) == 1
    passport_data = passport_module.load_passport(storage.passport_path(tmp_path, slug))
    assert passport_data["building_class"] == "Бизнес"
    assert passport_data["address"] == "вручную вписанный адрес"


def test_run_master_import_rejects_wrong_extension(tmp_path):
    app = create_app(tmp_path)
    client = app.test_client()

    resp = client.post(
        "/master-import",
        data={"workbook_file": (io.BytesIO(b"not excel"), "notes.txt")},
        content_type="multipart/form-data",
    )

    assert resp.status_code == 400
    assert "xlsx" in resp.get_data(as_text=True)


def test_project_page_does_not_show_the_raw_imported_table(tmp_path):
    # The "Импортировано из сводного файла" accordion (the full cost table
    # as the workbook has it) was removed from the page — it fed into no
    # calculation, so it was just a second, disconnected-looking cost table
    # sitting next to the real ones. The raw data is still saved to disk
    # (master_import.json), just not shown here.
    app = create_app(tmp_path)
    client = app.test_client()
    data = _workbook_bytes(
        "Nicole 1", None, "Жилая недвижимость (Делюкс)", "АО ФОДД",
        (10.0, 20.0, 30.0),
    )
    client.post(
        "/master-import",
        data={"workbook_file": (io.BytesIO(data), "portfolio.xlsx")},
        content_type="multipart/form-data",
    )
    [slug] = storage.list_project_slugs(tmp_path)

    resp = client.get(f"/projects/{slug}")
    body = resp.get_data(as_text=True)

    assert resp.status_code == 200
    assert "Импортировано из сводного файла" not in body
    assert storage.master_import_path(tmp_path, slug).exists()
