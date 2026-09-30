import io
from decimal import Decimal

import openpyxl
import pytest
from openpyxl import Workbook

from app import cost_increase, kp_analysis, predicted_increase

# Колонки маленькой книги той же формы, что и настоящая сводная тендерная
# таблица: № раздела, статья, наименование — и два блока подрядчиков по
# шесть колонок (материалы, СМР, косвенные, всего, комментарии, ожидаемая).
NUMBER_COL, ARTICLE_COL, NAME_COL = 2, 3, 4
BLOCKS = {"ООО «Альфа»": 6, "АО «Бета»": 12}


def _offer_workbook(rows):
    """Книга: шапка в строках 7/12/13, данные с 14-й.

    ``rows`` — ``(номер, название, {подрядчик: (материалы, смр, всего)})``.
    """
    wb = Workbook()
    ws = wb.active
    ws.cell(7, 5, "Наименование контрагента")
    ws.cell(12, NUMBER_COL, "№ раздела")
    ws.cell(12, ARTICLE_COL, "Статья СМР")
    ws.cell(12, NAME_COL, "Наименование работ")
    for name, start in BLOCKS.items():
        ws.cell(7, start, name)
        ws.cell(12, start, "Стоимость всего, RUB, ОСН, с учетом НДС 22%")
        for offset, label in enumerate(("Материалы", "СМР", "Косвенные расходы", "Всего")):
            ws.cell(13, start + offset, label)
        ws.cell(12, start + 4, "Комментарии")
        ws.cell(12, start + 5, "Ожидаемая стоимость")
    ws.cell(14, NUMBER_COL, "1")
    ws.cell(14, NAME_COL, "Лот №1 — генподряд")
    for index, (number, name, amounts) in enumerate(rows):
        row = 15 + index
        ws.cell(row, NUMBER_COL, number)
        ws.cell(row, ARTICLE_COL, name)
        for contractor, values in amounts.items():
            start = BLOCKS[contractor]
            materials, smr, total = values
            ws.cell(row, start, materials)
            ws.cell(row, start + 1, smr)
            ws.cell(row, start + 3, total)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


FACADE_ROWS = [
    ("6", "6. Фасадные работы", {
        "ООО «Альфа»": (700, 400, 1100),
        "АО «Бета»": (900, 500, 1400),
    }),
    ("6.1", "6.1. Навесной фасад", {
        "ООО «Альфа»": (700, 400, 1100),
        "АО «Бета»": (900, 500, 1400),
    }),
    ("6.2", "6.2. Витражи", {
        "ООО «Альфа»": (0, 0, 0),
        "АО «Бета»": (None, None, None),
    }),
    (None, "Монтаж витражей", {"ООО «Альфа»": (0, 0, 0)}),
    ("7", "7. Кровля", {}),
]


def test_parse_offer_finds_contractors_and_their_columns():
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(FACADE_ROWS)))

    assert [c.name for c in offer.contractors] == ["ООО «Альфа»", "АО «Бета»"]
    alpha = offer.contractors[0]
    assert (alpha.materials_col, alpha.smr_col, alpha.total_col) == (6, 7, 9)
    assert (alpha.comment_col, alpha.expected_col) == (10, 11)


def test_parse_offer_reads_sections_and_subsections_but_not_items_or_the_lot():
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(FACADE_ROWS)))

    assert [(line.number, line.level) for line in offer.lines] == [
        ("6", kp_analysis.LEVEL_SECTION),
        ("6.1", kp_analysis.LEVEL_SUBSECTION),
        ("6.2", kp_analysis.LEVEL_SUBSECTION),
        ("7", kp_analysis.LEVEL_SECTION),
    ]
    assert [line.key for line in offer.lines if line.level == kp_analysis.LEVEL_SECTION] == [
        "facade", "roof",
    ]
    facade = offer.lines[0]
    assert facade.amounts[0] == (Decimal("700"), Decimal("400"), Decimal("1100"))


def test_parse_offer_without_a_comments_column_is_an_error():
    wb = Workbook()
    wb.active.cell(1, 1, "что-то другое")
    buf = io.BytesIO()
    wb.save(buf)

    with pytest.raises(kp_analysis.OfferError):
        kp_analysis.parse_offer(io.BytesIO(buf.getvalue()))


def test_parse_offer_with_no_amounts_at_all_asks_to_resave_in_excel():
    rows = [("6", "6. Фасадные работы", {})]

    with pytest.raises(kp_analysis.OfferError, match="Excel"):
        kp_analysis.parse_offer(io.BytesIO(_offer_workbook(rows)))


def _project(name, building_class, area, estimate, vat="22%", signed=None, predicted=None):
    return kp_analysis.Project(
        name=name,
        passport={"building_class": building_class, "total_area_sqm": area, "vat": vat},
        estimate={key: Decimal(str(v)) for key, v in estimate.items()},
        signed_report=signed,
        predicted_report=predicted,
    )


def test_class_averages_add_signed_and_predicted_increases_to_the_estimate():
    signed = cost_increase.build_report(
        [cost_increase.Line("Фасадные работы", 1000, 1200)], {"facade": Decimal("1000")},
    )
    predicted = predicted_increase.build_report(
        [predicted_increase.Line("Фасадные работы", 100)],
    )
    projects = [
        _project("А", "Бизнес", 10, {"facade": 1000}, signed=signed, predicted=predicted),
        _project("Б", "Бизнес", 10, {"facade": 700}),
        _project("В", "Комфорт", 10, {"facade": 99999}),
    ]

    averages = kp_analysis.class_averages(projects, "Бизнес")

    # (1000 + 200 + 100 + 700) / (10 + 10)
    assert averages.per_sqm["facade"] == pytest.approx(100.0)
    assert averages.counts["facade"] == 2
    assert averages.considered == 2


def test_class_averages_bring_vat_to_22_percent_and_drop_unknown_vat():
    projects = [
        _project("А", "Бизнес", 10, {"facade": 1200}, vat="20%"),
        _project("Б", "Бизнес", 10, {"facade": 500}, vat="непонятно"),
    ]

    averages = kp_analysis.class_averages(projects, "Бизнес")

    assert averages.per_sqm["facade"] == pytest.approx(1200 * 122 / 120 / 10)
    assert averages.excluded == ["Б"]


def test_class_averages_skip_projects_without_estimate_or_area():
    projects = [
        _project("А", "Бизнес", None, {"facade": 1200}),
        _project("Б", "Бизнес", 10, {}),
    ]

    averages = kp_analysis.class_averages(projects, "Бизнес")

    assert averages.per_sqm == {}
    assert averages.considered == 0


def _averages(per_sqm):
    return kp_analysis.ClassAverages(
        per_sqm=per_sqm, counts={key: 1 for key in per_sqm}, considered=1, excluded=[],
    )


def _remarks_by_cell(analysis):
    return {(r.row, r.col): r.value for r in analysis.remarks}


def test_analyze_marks_sections_over_the_class_average():
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(FACADE_ROWS)))

    # Ожидаемая 1000 на 10 м²: Альфа +10%, Бета +40%.
    analysis = kp_analysis.analyze(offer, _averages({"facade": 100.0}), 10)
    remarks = _remarks_by_cell(analysis)

    assert remarks[(15, 10)] == kp_analysis.REMARK_OVERPRICED
    assert remarks[(15, 16)] == kp_analysis.REMARK_HEAVILY_OVERPRICED
    assert remarks[(15, 11)] == pytest.approx(1000.0)
    assert remarks[(15, 17)] == pytest.approx(1000.0)
    [facade] = analysis.sections
    assert [cell.deviation_pct for cell in facade.cells] == [
        pytest.approx(10.0), pytest.approx(40.0),
    ]
    assert facade.cells[0].per_sqm == pytest.approx(110.0)


def test_analyze_says_nothing_about_offers_at_or_below_the_average():
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(FACADE_ROWS)))

    analysis = kp_analysis.analyze(offer, _averages({"facade": 200.0}), 10)

    assert (15, 10) not in _remarks_by_cell(analysis)
    assert (15, 16) not in _remarks_by_cell(analysis)


def test_analyze_asks_to_price_empty_subsections():
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(FACADE_ROWS)))

    remarks = _remarks_by_cell(kp_analysis.analyze(offer, _averages({"facade": 100.0}), 10))

    assert remarks[(17, 10)] == kp_analysis.REMARK_PRICE_SUBSECTION
    assert remarks[(17, 16)] == kp_analysis.REMARK_PRICE_SUBSECTION


def test_analyze_treats_a_token_price_as_not_priced():
    rows = [
        ("6", "6. Фасадные работы", {"ООО «Альфа»": (70000, 30000, 100000)}),
        ("6.1", "6.1. Навесной фасад", {"ООО «Альфа»": (0, 0.01, 0.01)}),
    ]
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(rows)))

    remarks = _remarks_by_cell(kp_analysis.analyze(offer, _averages({}), 10))

    assert remarks[(16, 10)] == kp_analysis.REMARK_PRICE_SUBSECTION


def test_analyze_flags_smr_over_half_of_materials():
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(FACADE_ROWS)))

    remarks = _remarks_by_cell(kp_analysis.analyze(offer, _averages({"facade": 100.0}), 10))

    # 6.1: Альфа 400 из 700 (57%), Бета 500 из 900 (56%) — оба больше половины.
    assert remarks[(16, 10)] == kp_analysis.REMARK_UNJUSTIFIED_SMR
    assert remarks[(16, 16)] == kp_analysis.REMARK_UNJUSTIFIED_SMR


def test_analyze_does_not_flag_smr_at_or_below_half_of_materials():
    rows = [
        ("6", "6. Фасадные работы", {"ООО «Альфа»": (1000, 500, 1500)}),
        ("6.1", "6.1. Навесной фасад", {"ООО «Альфа»": (1000, 500, 1500)}),
    ]
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(rows)))

    remarks = _remarks_by_cell(kp_analysis.analyze(offer, _averages({}), 10))

    assert (16, 10) not in remarks


def test_analyze_flags_smr_with_no_materials_at_all():
    rows = [
        ("6", "6. Фасадные работы", {"ООО «Альфа»": (0, 3000, 3000)}),
        ("6.1", "6.1. Навесной фасад", {"ООО «Альфа»": (0, 3000, 3000)}),
    ]
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(rows)))

    remarks = _remarks_by_cell(kp_analysis.analyze(offer, _averages({}), 10))

    assert remarks[(16, 10)] == kp_analysis.REMARK_UNJUSTIFIED_SMR


def test_analyze_skips_sections_nobody_priced():
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(FACADE_ROWS)))

    analysis = kp_analysis.analyze(offer, _averages({"facade": 100.0, "roof": 50.0}), 10)

    assert [s.key for s in analysis.sections] == ["facade"]
    assert not any(r.row == 19 for r in analysis.remarks)


def test_analyze_asks_to_price_a_section_one_contractor_left_empty():
    rows = [
        ("6", "6. Фасадные работы", {"ООО «Альфа»": (700, 300, 1000), "АО «Бета»": (0, 0, 0)}),
    ]
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(rows)))

    remarks = _remarks_by_cell(kp_analysis.analyze(offer, _averages({"facade": 100.0}), 10))

    assert remarks[(15, 16)] == kp_analysis.REMARK_PRICE_SECTION


def test_analyze_leaves_a_section_without_class_average_unrated():
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(FACADE_ROWS)))

    analysis = kp_analysis.analyze(offer, _averages({}), 10)

    [facade] = analysis.sections
    assert facade.avg_per_sqm is None
    assert all(cell.deviation_pct is None for cell in facade.cells)
    assert (15, 10) not in _remarks_by_cell(analysis)
    assert (15, 11) not in _remarks_by_cell(analysis)
    # Подразделы оцениваются и без средней.
    assert _remarks_by_cell(analysis)[(17, 10)] == kp_analysis.REMARK_PRICE_SUBSECTION


def test_write_remarks_fills_empty_cells_and_keeps_existing_text():
    source = _offer_workbook(FACADE_ROWS)
    wb = openpyxl.load_workbook(io.BytesIO(source))
    wb.active.cell(15, 16, "уже написано вручную")
    buf = io.BytesIO()
    wb.save(buf)

    result = kp_analysis.write_remarks(buf.getvalue(), [
        kp_analysis.Remark(15, 10, "Завышена стоимость за раздел"),
        kp_analysis.Remark(15, 16, "Существенно завышена стоимость за раздел"),
        kp_analysis.Remark(15, 11, 1000.0),
    ])

    ws = openpyxl.load_workbook(io.BytesIO(result)).active
    assert ws.cell(15, 10).value == "Завышена стоимость за раздел"
    assert ws.cell(15, 16).value == "уже написано вручную"
    assert ws.cell(15, 11).value == 1000.0


def test_rank_orders_offers_from_cheapest_to_most_expensive():
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(FACADE_ROWS)))
    analysis = kp_analysis.analyze(offer, _averages({"facade": 100.0}), 10)

    ranking = kp_analysis.rank(analysis, 10)

    assert [(r.place, r.name) for r in ranking.overall] == [
        (1, "ООО «Альфа»"), (2, "АО «Бета»"),
    ]
    alpha, beta = ranking.overall
    assert alpha.deviation_pct == pytest.approx(10.0)
    assert ranking.expected_total == pytest.approx(1000.0)
    [facade] = ranking.sections
    assert [r.name for r in facade.rows] == ["ООО «Альфа»", "АО «Бета»"]


def test_rank_puts_an_unpriced_offer_last_without_a_place():
    rows = [
        ("6", "6. Фасадные работы", {"ООО «Альфа»": (70000, 30000, 100000), "АО «Бета»": (0, 0, 0)}),
    ]
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(rows)))
    analysis = kp_analysis.analyze(offer, _averages({}), 10)

    ranking = kp_analysis.rank(analysis, 10)

    [section] = ranking.sections
    assert [(r.place, r.name) for r in section.rows] == [(1, "ООО «Альфа»"), (None, "АО «Бета»")]
    assert ranking.overall[-1].name == "АО «Бета»"
    assert ranking.overall[-1].unpriced == 1


def test_rank_counts_sections_a_contractor_left_unpriced():
    rows = [
        ("6", "6. Фасадные работы", {"ООО «Альфа»": (7000, 3000, 10000), "АО «Бета»": (7000, 3000, 12000)}),
        ("7", "7. Кровля", {"ООО «Альфа»": (0, 0, 0), "АО «Бета»": (7000, 3000, 10000)}),
    ]
    offer = kp_analysis.parse_offer(io.BytesIO(_offer_workbook(rows)))
    ranking = kp_analysis.rank(kp_analysis.analyze(offer, _averages({}), 10), 10)

    alpha = next(r for r in ranking.overall if r.name == "ООО «Альфа»")
    assert alpha.unpriced == 1
    assert alpha.total == Decimal("10000")


def test_normalize_contractor_ignores_legal_form_quotes_and_case():
    assert kp_analysis.normalize_contractor('АО "ФОДД"') == kp_analysis.normalize_contractor("ФОДД")
    assert kp_analysis.normalize_contractor("ООО «Бюро Констракшн»") == "бюро констракшн"
    assert kp_analysis.normalize_contractor(None) == ""


def test_contractor_history_counts_our_objects_overall_and_in_the_class():
    passports = [
        {"project_name": "Никель 1", "general_contractor": "АО ФОДД",
         "building_class": "Делюкс", "contract_price_rub": 300.0},
        {"project_name": "Никель 2", "general_contractor": "ФОДД",
         "building_class": "Бизнес", "contract_price_rub": 100.0},
        {"project_name": "Без цены", "general_contractor": "фодд",
         "building_class": "Делюкс", "contract_price_rub": None},
        {"project_name": "Чужой", "general_contractor": "АНТТЕК",
         "building_class": "Делюкс", "contract_price_rub": 999.0},
    ]

    history = kp_analysis.contractor_history(['АО "ФОДД"', "ООО «Новичок»"], passports, "Делюкс")

    fodd = history['АО "ФОДД"']
    assert (fodd.count, fodd.total) == (3, 400.0)
    assert (fodd.class_count, fodd.class_total) == (2, 300.0)
    assert [o.name for o in fodd.objects] == ["Никель 1", "Без цены", "Никель 2"]
    assert history["ООО «Новичок»"].count == 0
