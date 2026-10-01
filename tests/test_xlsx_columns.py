import io

import openpyxl
from openpyxl import Workbook

from app import xlsx_columns


def _roundtrip(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return openpyxl.load_workbook(io.BytesIO(buf.getvalue()))


def test_formulas_point_at_the_same_cells_after_the_insert():
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws["B1"], ws["C1"], ws["D1"] = 1, 2, 3
    ws["A2"] = "=B1+C1+$D$1"
    ws["E2"] = "=SUM(B1:D1)+LOG10(D1)"
    ws["F2"] = "=SUM(C:C)+Лист1!D1+'Другой лист'!D1"

    xlsx_columns.insert_blank_columns(ws, [3])

    assert ws["A2"].value == "=B1+D1+$E$1"
    assert ws["F2"].value == "=SUM(B1:E1)+LOG10(E1)"
    assert ws["G2"].value == "=SUM(D:D)+Лист1!E1+'Другой лист'!D1"
    assert ws["C1"].value is None and ws["D1"].value == 2


def test_merged_ranges_widths_groups_and_filter_follow_the_columns():
    wb = Workbook()
    ws = wb.active
    ws["D1"] = "Подрядчик"
    ws.merge_cells("D1:F1")
    ws.merge_cells("A2:E2")
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions.group("F", "G", hidden=True, outline_level=1)
    ws.auto_filter.ref = "A3:F10"

    xlsx_columns.insert_blank_columns(ws, [3], width=2.5)
    ws = _roundtrip(wb).active

    merged = {str(r) for r in ws.merged_cells.ranges}
    assert merged == {"E1:G1", "A2:F2"}
    assert ws["E1"].value == "Подрядчик"
    assert ws.column_dimensions["F"].width == 30
    assert ws.column_dimensions["C"].width == 2.5
    assert ws.column_dimensions["G"].hidden and ws.column_dimensions["G"].outlineLevel == 1
    assert ws.auto_filter.ref == "A3:G10"


def test_several_columns_are_inserted_in_one_pass():
    wb = Workbook()
    ws = wb.active
    for col, value in zip("ABCDE", range(1, 6)):
        ws[f"{col}1"] = value
    ws["A2"] = "=A1+B1+C1+D1+E1"
    ws.merge_cells("B3:D3")

    xlsx_columns.insert_blank_columns(ws, [2, 4])

    assert [ws.cell(1, c).value for c in range(1, 8)] == [1, None, 2, 3, None, 4, 5]
    assert ws["A2"].value == "=A1+C1+D1+F1+G1"
    assert {str(r) for r in ws.merged_cells.ranges} == {"C3:F3"}


def test_formatting_at_the_sheet_edge_does_not_spill_past_the_last_column():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.column_dimensions.group("XDG", "XFD", hidden=True)
    ws.cell(1, 16384).number_format = "0.00"

    xlsx_columns.insert_blank_columns(ws, [2])

    assert max(dim.max or 0 for dim in ws.column_dimensions.values()) <= 16384
    assert max(col for _row, col in ws._cells) <= 16384
    _roundtrip(wb)
