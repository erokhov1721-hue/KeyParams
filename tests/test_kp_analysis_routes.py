import io
import re
from decimal import Decimal

import openpyxl

from app import create_app, kp_analysis, passport as passport_module, routes, storage
from tests.test_kp_analysis import FACADE_ROWS, _offer_workbook


def _post(client, data=None, building_class="Бизнес", area="10", filename="КП.xlsx"):
    form = {"building_class": building_class, "area": area}
    if data is not None:
        form["offer_file"] = (io.BytesIO(data), filename)
    return client.post("/kp-analysis", data=form, content_type="multipart/form-data")


def _make_project(root, name, building_class, area, vat="22%"):
    slug = storage.create_project(root, name)
    passport_module.save_passport(
        {"project_name": name, "building_class": building_class,
         "total_area_sqm": area, "vat": vat, "ocr_fields": []},
        storage.passport_path(root, slug),
    )
    return slug


def test_kp_analysis_requires_a_file(tmp_path):
    client = create_app(tmp_path).test_client()

    resp = _post(client)

    assert resp.status_code == 400
    assert "Выберите файл" in resp.get_data(as_text=True)


def test_kp_analysis_requires_a_building_class(tmp_path):
    client = create_app(tmp_path).test_client()

    resp = _post(client, _offer_workbook(FACADE_ROWS), building_class="")

    assert resp.status_code == 400
    assert "класс жилья" in resp.get_data(as_text=True)


def test_kp_analysis_requires_a_positive_area(tmp_path):
    client = create_app(tmp_path).test_client()

    resp = _post(client, _offer_workbook(FACADE_ROWS), area="ноль")

    assert resp.status_code == 400
    assert "площадь" in resp.get_data(as_text=True)


def test_kp_analysis_rejects_a_workbook_that_is_not_an_offer_table(tmp_path):
    client = create_app(tmp_path).test_client()
    wb = openpyxl.Workbook()
    wb.active.cell(1, 1, "просто таблица")
    buf = io.BytesIO()
    wb.save(buf)

    resp = _post(client, buf.getvalue())

    assert resp.status_code == 400
    assert "Комментарии" in resp.get_data(as_text=True)


def test_kp_analysis_compares_against_projects_of_the_chosen_class(tmp_path, monkeypatch):
    client = create_app(tmp_path).test_client()
    _make_project(tmp_path, "Бизнес-объект", "Бизнес", 10)
    _make_project(tmp_path, "Комфорт-объект", "Комфорт", 10)
    estimates = {
        "Бизнес-объект": {"facade": Decimal("1000")},
        "Комфорт-объект": {"facade": Decimal("99999")},
    }
    monkeypatch.setattr(routes, "_estimate_totals", lambda root, slug: estimates[slug])

    resp = _post(client, _offer_workbook(FACADE_ROWS))
    body = resp.get_data(as_text=True)

    assert resp.status_code == 200
    assert "ООО «Альфа»" in body
    # Средняя 100 ₽/м² по одному объекту «Бизнес», Альфа +10%, Бета +40%.
    assert "+10,0 %" in body
    assert "+40,0 %" in body
    assert "Скачать отработанный файл" in body


def test_kp_analysis_download_returns_the_filled_workbook(tmp_path, monkeypatch):
    client = create_app(tmp_path).test_client()
    _make_project(tmp_path, "Бизнес-объект", "Бизнес", 10)
    monkeypatch.setattr(routes, "_estimate_totals", lambda root, slug: {"facade": Decimal("1000")})

    body = _post(client, _offer_workbook(FACADE_ROWS)).get_data(as_text=True)
    [link] = re.findall(r'href="(/kp-analysis/download/[0-9a-f]{32})"', body)
    resp = client.get(link)

    assert resp.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(resp.data)).active
    assert ws.cell(15, 10).value == kp_analysis.REMARK_OVERPRICED
    assert ws.cell(15, 16).value == kp_analysis.REMARK_HEAVILY_OVERPRICED
    assert ws.cell(15, 11).value == 1000.0


def test_kp_analysis_results_live_outside_the_projects_folder(tmp_path, monkeypatch):
    client = create_app(tmp_path).test_client()

    _post(client, _offer_workbook(FACADE_ROWS))

    assert storage.list_project_slugs(tmp_path) == []
    assert list(storage.kp_analysis_dir(tmp_path).glob("*.xlsx"))


def test_kp_analysis_download_of_an_unknown_token_is_not_found(tmp_path):
    client = create_app(tmp_path).test_client()

    assert client.get("/kp-analysis/download/" + "0" * 32).status_code == 404
    assert client.get("/kp-analysis/download/..%2Fsecret_key").status_code == 404


def test_kp_analysis_offers_the_ranking_as_a_pdf(tmp_path, monkeypatch):
    import pdfplumber

    client = create_app(tmp_path).test_client()
    _make_project(tmp_path, "Бизнес-объект", "Бизнес", 10)
    monkeypatch.setattr(routes, "_estimate_totals", lambda root, slug: {"facade": Decimal("1000")})

    body = _post(client, _offer_workbook(FACADE_ROWS)).get_data(as_text=True)
    [link] = re.findall(r'href="(/kp-analysis/download/[0-9a-f]{32}/pdf)"', body)
    resp = client.get(link)

    assert resp.status_code == 200
    assert resp.mimetype == "application/pdf"
    with pdfplumber.open(io.BytesIO(resp.data)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "рейтинг предложений" in text.lower()
    assert "ООО «Альфа»" in text
    assert "+40,0 %" in text
    assert "Фасадные работы" in text


def test_kp_ranking_pdf_of_an_unknown_token_is_not_found(tmp_path):
    client = create_app(tmp_path).test_client()

    assert client.get("/kp-analysis/download/" + "0" * 32 + "/pdf").status_code == 404


def test_kp_analysis_shows_the_contractors_objects_with_us(tmp_path, monkeypatch):
    client = create_app(tmp_path).test_client()
    for name, gc, cls, price in (
        ("Наш объект 1", "Альфа", "Бизнес", 2_000_000_000.0),
        ("Наш объект 2", "ООО Альфа", "Комфорт", 500_000_000.0),
    ):
        slug = storage.create_project(tmp_path, name)
        passport_module.save_passport(
            {"project_name": name, "general_contractor": gc, "building_class": cls,
             "contract_price_rub": price, "ocr_fields": []},
            storage.passport_path(tmp_path, slug),
        )
    monkeypatch.setattr(routes, "_estimate_totals", lambda root, slug: {})

    body = _post(client, _offer_workbook(FACADE_ROWS)).get_data(as_text=True)

    assert "Объекты MR Group" in body
    assert "Наш объект 1" in body and "Наш объект 2" in body
    assert "2 объект(ов) на 2,5 млрд ₽" in body
    assert "в классе «Бизнес» — 1 на 2,0 млрд ₽" in body
    assert "у нас объектов нет" in body  # АО «Бета»
    assert "К лучшему" not in body


def test_kp_download_button_is_disabled_until_an_analysis_is_run(tmp_path):
    client = create_app(tmp_path).test_client()

    body = client.get("/kp-analysis").get_data(as_text=True)

    assert re.search(r'<button[^>]*id="kp-download-xlsx"[^>]*disabled', body)
    assert "/kp-analysis/download/" not in body
