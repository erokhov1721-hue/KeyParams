"""Reading a project's predicted cost-increase workbook — one workbook per
project, its own "Предполагаемое ДС" figure per kind of work.

Unlike the cost-increase workbook (see ``cost_increase.py``), there is no
"было"/"стало" pair to compare: this file already states the predicted
increase itself, one number per kind of work, so reading it is a matter of
finding that one column and adding its rows up by section — no baseline.

Rows are named the way the estimate names its sections, so which report line
a row belongs to is decided by ``estimate_sections.classify``, the same
judgement ``cost_increase`` makes, which is what lets a figure here and a
figure in the estimate be about the same kind of work.
"""

import logging
import re
from collections import namedtuple
from decimal import Decimal

import openpyxl

from . import estimate_sections, extractors

logger = logging.getLogger(__name__)

# How far into the sheet to look for the header. The workbook opens straight
# into its table, but a title row or two above it costs nothing to allow for.
HEADER_SEARCH_ROWS = 20
HEADER_SEARCH_COLS = 40

# Matched as a substring, so both the correct spelling ("предполагаемое") and
# the typo actually seen in the field ("предпологаемое") are found — they
# share this stem.
AMOUNT_HEADER = "предпол"

# The workbook's own bottom line, read past rather than through — same rule
# as cost_increase.py's TOTAL_ROW_RE.
TOTAL_ROW_RE = re.compile(r"^(итого|всего|сумма|total)\b")
_LETTER_RE = re.compile(r"[a-zа-яё]", re.IGNORECASE)


class PredictedIncreaseError(Exception):
    """The file isn't a predicted-increase workbook this reader can make
    sense of."""


PROBLEM_MESSAGES = {
    "format": "Файл прогнозируемого удорожания должен быть в формате .xlsx",
    "too_big": "Файл прогнозируемого удорожания слишком большой — до 15 МБ",
    "unreadable": (
        "Не удалось прочитать файл — нужна таблица со столбцом "
        "«Предполагаемое ДС» по видам работ. Прежний файл оставлен на месте."
    ),
}

# One row of the workbook as it is written there.
Line = namedtuple("Line", "name amount")

# One line of the report: ``sources`` are the workbook's own row names added
# together into this line, carried so the placement can be checked rather
# than taken on trust — same idea as cost_increase.Row.sources.
Row = namedtuple("Row", "key label sources amount")

# ``rows`` in the report's own order, ``total`` their sum as a row of the
# same shape, ``unmatched`` the workbook's row names that belong to no
# report line.
Report = namedtuple("Report", "rows total unmatched")


def _text(ws, row, col):
    value = ws.cell(row=row, column=col).value
    return str(value).strip().lower() if value is not None else ""


def _named(value):
    """The value as a name, or None if it holds no letters — a row number, a
    ledger code, an empty cell."""
    text = str(value or "").strip()
    return text if _LETTER_RE.search(text) else None


def _amount(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    return extractors.parse_money(str(value))


Header = namedtuple("Header", "row name_col amount_col")


def _find_amount_col(ws, row):
    for col in range(1, HEADER_SEARCH_COLS + 1):
        if AMOUNT_HEADER in _text(ws, row, col):
            return col
    return None


def _name_column(ws, header_row, amount_col):
    """The column the kinds of work are named in — same left-of-the-money-
    first search as cost_increase._name_column."""
    def count(col):
        return sum(
            1 for row in range(header_row + 1, ws.max_row + 1)
            if _named(ws.cell(row=row, column=col).value)
        )

    for candidates in (
        range(amount_col - 1, 0, -1),
        (c for c in range(1, HEADER_SEARCH_COLS + 1) if c != amount_col),
    ):
        counts = [(count(col), col) for col in candidates]
        best = max(counts, default=(0, None), key=lambda pair: pair[0])
        if best[0]:
            return best[1]
    return None


def _find_header(ws):
    for row in range(1, HEADER_SEARCH_ROWS + 1):
        amount_col = _find_amount_col(ws, row)
        if amount_col is None:
            continue
        name_col = _name_column(ws, row, amount_col)
        if name_col is not None:
            return Header(row, name_col, amount_col)
    return None


def read_lines(source) -> list:
    """Every priced row of a predicted-increase workbook, as it is written
    there.

    ``source`` is a path or an open file — the upload is checked before it
    is saved, so a file that turns out to be unreadable can be refused
    without having overwritten the one that was already there.
    """
    try:
        wb = openpyxl.load_workbook(source, data_only=True)
    except Exception as e:
        raise PredictedIncreaseError(f"файл не читается как .xlsx: {e}") from e

    for ws in wb.worksheets:
        header = _find_header(ws)
        if header is None:
            continue
        lines = _lines_from_sheet(ws, header)
        if lines:
            return lines

    raise PredictedIncreaseError(
        'в файле нет таблицы со столбцом «Предполагаемое ДС»'
    )


def _lines_from_sheet(ws, header) -> list:
    """Every priced row read on its own account, numbered hierarchy or not.

    Tempting as it looks next to a smeta's own "укрупненная" shape (see
    estimate_sections._sections_from_flat_numbered), this file's numbering
    is not a rollup: a real workbook's own grand-total cell is a plain
    ``SUBTOTAL(9, ...)`` over every row in the range, top-level and
    sub-level alike, which only adds up if each row is its own independent
    predicted change rather than a parent already including its children.
    A row numbered "5.5" under "5" is a specific extra adjustment noted
    against that one sub-item, on top of "5"'s own figure, not a part of it
    restated.
    """
    lines = []
    for row in range(header.row + 1, ws.max_row + 1):
        name = _named(ws.cell(row=row, column=header.name_col).value)
        if not name or TOTAL_ROW_RE.match(name.lower()):
            continue
        amount = _amount(ws.cell(row=row, column=header.amount_col).value)
        if amount is None:
            continue
        lines.append(Line(name, amount))
    return lines


def build_report(lines) -> Report:
    """The workbook's rows added up by kind of work.

    Lines are returned in the report's own order rather than the workbook's,
    so the kinds of work read down the accordion in the order they read down
    the estimate and the cost-increase report.
    """
    gathered, unmatched = {}, []
    for line in lines:
        key = estimate_sections.classify(line.name)
        if key is None:
            unmatched.append(line.name)
            continue
        amount, sources = gathered.get(key, (Decimal("0"), []))
        gathered[key] = (amount + line.amount, sources + [line.name])

    rows = []
    for key in estimate_sections.CATEGORY_KEYS:
        if key not in gathered:
            continue
        amount, sources = gathered[key]
        if not amount:
            continue
        rows.append(Row(
            key=key, label=estimate_sections.CATEGORY_LABELS.get(key, key),
            sources=sources, amount=amount,
        ))

    if unmatched:
        # Not an error: the workbook may carry a kind of work this report has
        # no line for. Named on the page too, so a row quietly missing from
        # the table can be traced to the row that was skipped.
        logger.info(
            "Разделы файла прогнозируемого удорожания без строки в отчёте: %s",
            "; ".join(unmatched),
        )

    total = Row(
        key=None, label="Итого", sources=[],
        amount=sum((row.amount for row in rows), Decimal("0")),
    )
    return Report(rows=rows, total=total, unmatched=unmatched)


def read_report(source) -> Report:
    return build_report(read_lines(source))
