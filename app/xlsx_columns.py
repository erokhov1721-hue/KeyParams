"""Вставка пустых колонок в лист openpyxl так, как это делает Excel.

``Worksheet.insert_cols`` у openpyxl двигает только сами ячейки, и медленно:
на тендерной таблице с двумястами тысячами оформленных ячеек — минуты на
одну колонку. Ссылки в формулах, объединённые области, ширины и
группировки колонок и автофильтр он оставляет на старых местах — после него
«=AC17» в сдвинутой ячейке смотрит уже не туда, а итоги таблицы считают
чужие колонки. Здесь всё это сдвигается за один проход, сразу для всех
вставляемых колонок.
"""

import bisect
import re

from openpyxl.cell.cell import MergedCell
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.dimensions import ColumnDimension

# Последняя колонка листа Excel (XFD). Оформление в настоящих файлах часто
# дотянуто до неё самой, и после сдвига вправо ушло бы за край — такой файл
# Excel не открывает.
MAX_COLUMN = 16384

_COL_RE = re.compile(r"(\$?)([A-Z]{1,3})(?=\$?\d|$)")
_A1_RANGE_RE = re.compile(r"^\$?[A-Z]{1,3}(\$?\d+)?(:\$?[A-Z]{1,3}(\$?\d+)?)?$")


class _Shift:
    """Куда уезжает колонка: ``points`` — номера колонок (в исходной
    нумерации), перед каждой из которых встаёт новая пустая."""

    def __init__(self, points):
        self.points = sorted(points)

    def __call__(self, col):
        return col + bisect.bisect_right(self.points, col)


def _shift_ref(ref, shift):
    """Ссылка в нотации A1 без имени листа — с колонками на новых местах."""
    def move(match):
        dollar, letters = match.group(1), match.group(2)
        return dollar + get_column_letter(shift(column_index_from_string(letters)))
    return ":".join(_COL_RE.sub(move, part, count=1) for part in ref.split(":"))


def _shift_formula(formula, shift, sheet_title):
    tokens = Tokenizer(formula).items
    changed = False
    for token in tokens:
        if token.type != Token.OPERAND or token.subtype != Token.RANGE:
            continue
        value = token.value
        sheet = None
        if "!" in value:
            sheet, value = value.rsplit("!", 1)
            if sheet.strip("'") != sheet_title:
                continue  # ссылка на другой лист — там колонки не вставляли
        if not _A1_RANGE_RE.match(value):
            continue  # имя диапазона, ссылка на строки «1:1» и т. п.
        shifted = _shift_ref(value, shift)
        if shifted != value:
            token.value = f"{sheet}!{shifted}" if sheet is not None else shifted
            changed = True
    if not changed:
        return formula
    return "=" + "".join(token.value for token in tokens)


def _shift_cells(ws, shift):
    moved = {}
    for (row, col), cell in ws._cells.items():
        new_col = shift(col)
        if new_col > MAX_COLUMN:
            if cell.value not in (None, ""):
                raise ValueError(
                    f"Ячейка {cell.coordinate} со значением ушла бы за последнюю колонку листа"
                )
            continue  # пустая ячейка с одним оформлением у самого края листа
        if new_col != col:
            cell.column = new_col
        moved[(row, new_col)] = cell
        value = cell.value
        if isinstance(value, str) and value.startswith("="):
            cell.value = _shift_formula(value, shift, ws.title)
    ws._cells = moved


def _shift_merged(ws, shift, points):
    """Объединённые области не разъединяются (так потерялись бы рамки их
    внутренних клеток): сдвигаются их границы; область, внутрь которой
    встала новая колонка, растёт на неё."""
    ranges = list(ws.merged_cells.ranges)
    ws.merged_cells = MultiCellRange()
    for rng in ranges:
        old_min, old_max = rng.min_col, rng.max_col
        new_min, new_max = shift(old_min), shift(old_max)
        rng.shift(col_shift=new_min - old_min)
        rng.expand(right=(new_max - new_min) - (old_max - old_min))
        for point in points:
            if old_min < point <= old_max:
                new_col = shift(point) - 1
                for row in range(rng.min_row, rng.max_row + 1):
                    ws._cells[(row, new_col)] = MergedCell(ws, row=row, column=new_col)
        ws.merged_cells.add(rng)


def _shift_dimensions(ws, shift, points, width):
    """Ширины, скрытые колонки и группировки — вслед за колонками; группа,
    внутрь которой встала новая колонка, делится вокруг неё."""
    pieces = []
    for dim in list(ws.column_dimensions.values()):
        # Ещё не сохранявшаяся колонка знает только свою букву, а не
        # диапазон min..max — тогда это диапазон из одной колонки.
        lo = dim.min or column_index_from_string(dim.index)
        hi = dim.max or lo
        start = lo
        for point in points:
            if lo < point <= hi:
                pieces.append((start, point - 1, dim))
                start = point
        pieces.append((start, hi, dim))
    ws.column_dimensions.clear()
    for lo, hi, dim in pieces:
        new_lo, new_hi = shift(lo), min(shift(hi), MAX_COLUMN)
        if new_lo > MAX_COLUMN:
            continue
        letter = get_column_letter(new_lo)
        new = ColumnDimension(
            ws, index=letter, width=dim.width, bestFit=dim.bestFit, hidden=dim.hidden,
            outlineLevel=dim.outlineLevel, collapsed=dim.collapsed,
            customWidth=dim.customWidth,
        )
        new._style = dim._style
        new.min, new.max = new_lo, new_hi
        ws.column_dimensions[letter] = new
    for point in points:
        letter = get_column_letter(shift(point) - 1)
        ws.column_dimensions[letter] = ColumnDimension(ws, index=letter, width=width)


def insert_blank_columns(ws, points, width=2.43):
    """Вставляет по пустой колонке перед каждой из колонок ``points`` (номера
    с единицы, в нумерации листа до вставки), сдвигая вправо всё правее:
    ячейки, ссылки в формулах этого листа, объединённые области, ширины и
    группировки колонок, автофильтр. Новые колонки — шириной ``width``."""
    points = sorted(set(points))
    if not points:
        return
    shift = _Shift(points)
    _shift_cells(ws, shift)
    _shift_merged(ws, shift, points)
    _shift_dimensions(ws, shift, points, width)
    if ws.auto_filter.ref:
        ws.auto_filter.ref = _shift_ref(ws.auto_filter.ref, shift)
