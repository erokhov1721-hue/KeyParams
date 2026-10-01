"""Анализ КП: сводная тендерная таблица подрядчиков против расчётной
стоимости MR Group и средней ₽/м² по загруженным объектам того же класса.

Файл — сводная оценочная таблица тендера: слева номер раздела, статья и
наименование работ, правее — блок «Расчетная стоимость» (стоимость MR Group)
и по блоку колонок на каждого подрядчика (стоимость всего: материалы / СМР /
косвенные / всего; иногда ещё «Комментарии» и «Ожидаемая стоимость»). Этот
модуль:

- ``parse_offer`` читает из неё блоки подрядчиков, блок расчётной стоимости
  и строки разделов и подразделов (позиции внутри подразделов не нужны —
  оценка идёт не по ним);
- ``class_averages`` считает среднюю ₽/м² по виду работ по объектам класса —
  смета плюс подписанное и прогнозируемое удорожание, всё при НДС 22%;
- ``analyze`` сопоставляет одно с другим и решает, что написать в
  «Комментарии» (и «Ожидаемую стоимость», если такая колонка есть) каждого
  подрядчика. Заполнена расчётная стоимость — она эталон, и сравнение идёт
  с ней на каждом разделе и подразделе; пуста — со средней по классу;
- ``write_remarks`` вписывает это в исходный файл, не трогая остального.
  Колонки «Комментарии» нет — замечания пишутся в пустую колонку сразу за
  блоком подрядчика, с тем же заголовком.

Ничего здесь не знает про Flask и не читает файлы проектов: объекты
собирает ``app.routes`` теми же функциями, что и «Сводка по удорожанию».
"""

import copy
import io
import math
import re
from collections import namedtuple
from decimal import Decimal

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from . import comparison, estimate_sections, xlsx_columns

# НДС, при котором подрядчики дают цены в тендерной таблице («с учетом НДС
# 22%»). Цены загруженных объектов приводятся к нему же, иначе сравнение
# шло бы между разными налоговыми базами.
OFFER_VAT_RATE = 22.0

# Пороги оценки — отклонение предложения от эталона (расчётной стоимости
# или ожидаемой по средней класса).
HEAVILY_OVERPRICED_PCT = 30.0

# СМР больше этой доли от стоимости материалов — «необоснована стоимость СМР».
SMR_TO_MATERIALS_LIMIT = Decimal("0.5")

# Сумма меньше этой — не цена, а заглушка: подрядчики ставят 0,01 ₽ или
# 90 ₽ там, где раздел не расценивали, и специалист помечает такие строки
# «расценить» так же, как пустые.
NOMINAL_PRICE_LIMIT = Decimal("1000")

REMARK_OVERPRICED = "Завышена стоимость за раздел"
REMARK_HEAVILY_OVERPRICED = "Существенно завышена стоимость за раздел"
REMARK_SUB_OVERPRICED = "Завышена стоимость"
REMARK_SUB_HEAVILY_OVERPRICED = "Существенно завышена стоимость"
REMARK_PRICE_SECTION = "расценить раздел."
REMARK_PRICE_SUBSECTION = "расценить подраздел."
REMARK_UNJUSTIFIED_SMR = "необоснована стоимость СМР"

COMMENTS_TITLE = "Комментарии"
# Ширина колонки замечаний, которую программа заводит сама: в таблице на
# этом месте узкий разделитель между блоками подрядчиков.
CREATED_COMMENT_WIDTH = 55
# Разделитель, который встаёт за этой колонкой вместо занятого ею, — той
# же ширины, что был в таблице.
SPACER_WIDTH = 2.43
# Сколько знаков замечания помещается в строку колонки такой ширины: ширина
# колонки Excel меряется в цифрах «0», а русские буквы и заглавные шире.
CHARS_PER_LINE = 45
# Высота строки текста — во столько раз больше кегля шрифта ячейки.
LINE_SPACING = 1.45

LEVEL_SECTION = "section"
LEVEL_SUBSECTION = "subsection"

HEADER_SEARCH_ROWS = 40
# Сколько строк ниже «Наименование контрагента» может занимать шапка.
HEADER_DEPTH = 10

COMMENTS_HEADER = "комментарии"
EXPECTED_HEADER = "ожидаем"
COST_HEADER = "стоимость всего"
CUSTOMER_VOLUMES_HEADER = "заказчик"
NAME_ROW_HEADER = "наименование контрагента"
REFERENCE_HEADER = "расчетная стоимость"
NUMBER_HEADER = "№ раздела"
ARTICLE_HEADER = "статья"
WORK_NAME_HEADER = "наименование работ"

SECTION_NUMBER_RE = re.compile(r"^\d+\.?$")
SUBSECTION_NUMBER_RE = re.compile(r"^\d+(\.\d+)+\.?$")
LOT_RE = re.compile(r"^лот\b", re.IGNORECASE)


class OfferError(Exception):
    """Файл — не сводная тендерная таблица, или в нём нечего сравнивать."""


# Колонки одного блока (номера с единицы, как в openpyxl). ``comment_col`` —
# куда писать замечания (None — писать некуда), ``comment_created`` — этой
# колонки в файле не было, её заголовок ставит программа в ``header_row``,
# а рамку тянет от ``top_row`` (строка с названиями подрядчиков) вниз.
Contractor = namedtuple(
    "Contractor",
    "name materials_col smr_col total_col comment_col expected_col "
    "header_row comment_created top_row",
)

# Строка раздела или подраздела. ``amounts`` — по подрядчику, в порядке
# ``Offer.contractors``: ``(материалы, смр, всего)``, каждое Decimal или None;
# ``reference`` — то же по расчётной стоимости (или None, если её блока нет).
Line = namedtuple("Line", "row number name level key amounts reference")

# ``reference`` — блок расчётной стоимости MR Group или None;
# ``has_reference`` — заполнена ли она хоть по одному разделу.
Offer = namedtuple("Offer", "contractors lines reference has_reference")

# Объект из базы: паспорт, смета по видам работ и два отчёта по удорожанию
# (любой может быть None) — в том виде, в каком их отдают ``app.routes``.
Project = namedtuple("Project", "name passport estimate signed_report predicted_report")

# ``per_sqm`` — {вид работ: ₽/м²}, ``counts`` — по скольким объектам,
# ``considered`` — сколько объектов класса вообще вошло в расчёт,
# ``excluded`` — имена объектов класса, выпавших из-за неизвестного НДС.
ClassAverages = namedtuple("ClassAverages", "per_sqm counts considered excluded")

# Что вписать в какую ячейку: текст замечания или ожидаемую стоимость.
Remark = namedtuple("Remark", "row col value")

# Одна клетка таблицы на странице: подрядчик × раздел. ``deviation_pct`` —
# от средней по классу, ``ref_deviation_pct`` — от расчётной стоимости.
Cell = namedtuple("Cell", "total per_sqm deviation_pct ref_deviation_pct remark")
Section = namedtuple(
    "Section",
    "row number name key label avg_per_sqm count expected reference cells",
)
# ``sections`` — разделы верхнего уровня (из них складывается общий итог),
# ``groups`` — то, по чему строится рейтинг «по видам работ»: те же разделы
# или, если раздел в файле всего один, его подразделы первого уровня.
# ``new_columns`` — колонки «Комментарии», которые программа заводит сама.
Analysis = namedtuple(
    "Analysis", "contractors sections groups remarks has_reference new_columns",
)


def _text(value):
    return " ".join(str(value).split()).lower() if value is not None else ""


def _amount(value):
    """Число из ячейки как Decimal, или None — пустая ячейка, текст,
    ошибка формулы вроде «#DIV/0!»."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    return None


def _find_col(ws, row, predicate, start=1, end=None):
    end = min(end or ws.max_column, ws.max_column)
    for col in range(start, end + 1):
        if predicate(_text(ws.cell(row, col).value)):
            return col
    return None


def _find_name_row(ws):
    for row in range(1, min(HEADER_SEARCH_ROWS, ws.max_row) + 1):
        col = _find_col(ws, row, lambda t: t == NAME_ROW_HEADER)
        if col:
            return row, col
    raise OfferError(
        "Не похоже на сводную тендерную таблицу: не нашлось строки "
        "«Наименование контрагента» с названиями подрядчиков."
    )


def _block_starts(ws, name_row, label_col):
    """``[(колонка, название)]`` — где в строке названий начинается каждый
    блок: расчётная стоимость и подрядчики."""
    starts = []
    for col in range(label_col + 1, ws.max_column + 1):
        value = ws.cell(name_row, col).value
        if value is not None and str(value).strip():
            starts.append((col, " ".join(str(value).split())))
    return starts


def _header_rows(ws, name_row):
    return range(name_row, min(name_row + HEADER_DEPTH, ws.max_row) + 1)


def _block_columns(ws, name_row, start, end):
    """``(материалы, смр, всего, строка шапки)`` блока, или None.

    «Стоимость всего» ищется в любой строке шапки — у расчётной стоимости и
    у подрядчиков она бывает на разных строках, — а подписи «Материалы» /
    «СМР» / «Всего» — строкой ниже, в пределах четырёх колонок.
    """
    for row in _header_rows(ws, name_row):
        for col in range(start, end + 1):
            text = _text(ws.cell(row, col).value)
            if not text.startswith(COST_HEADER) or CUSTOMER_VOLUMES_HEADER in text:
                continue
            sub_row = row + 1
            last = min(col + 4, end)

            def sub(label, col=col, last=last, sub_row=sub_row):
                return _find_col(ws, sub_row, lambda t: t == label, col, last)

            materials, smr, total = sub("материалы"), sub("смр"), sub("всего")
            if materials and smr and total:
                return materials, smr, total, row
    return None


def _column_is_empty(ws, col):
    return all(ws.cell(row, col).value in (None, "") for row in range(1, ws.max_row + 1))


def _labelled_col(ws, name_row, start, end, predicate):
    for row in _header_rows(ws, name_row):
        col = _find_col(ws, row, predicate, start, end)
        if col:
            return col
    return None


def _level(number):
    if SECTION_NUMBER_RE.match(number):
        return LEVEL_SECTION
    if SUBSECTION_NUMBER_RE.match(number):
        return LEVEL_SUBSECTION
    return None


def parse_offer(source) -> Offer:
    """Блоки подрядчиков, расчётная стоимость и строки разделов/подразделов
    первого листа.

    Суммы берутся такими, какими их последний раз посчитал Excel: формулы
    таблицы здесь не пересчитываются.
    """
    try:
        wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001 — любой нечитаемый файл это одна и та же ошибка
        raise OfferError("Не удалось открыть файл как таблицу Excel.") from e
    try:
        return _parse_sheet(_SheetView(wb.worksheets[0]))
    finally:
        wb.close()


class _SheetView:
    """Лист, прочитанный целиком в память: ``cell(row, col).value``, как у
    обычного листа openpyxl, но без его медленного доступа по ячейке в
    режиме read_only.

    Ширина берётся по шапке, а не по ``max_column`` листа: в настоящих файлах
    оформление тянется до колонки XFD, и читать шестнадцать тысяч пустых
    колонок на каждой из двух тысяч строк — это сотни мегабайт впустую.
    """

    # Запас справа от последней непустой колонки шапки — колонка-разделитель
    # за последним подрядчиком, куда могут лечь замечания.
    _MARGIN = 3

    class _Cell:
        __slots__ = ("value", "row")

        def __init__(self, value, row):
            self.value = value
            self.row = row

    def __init__(self, ws):
        head = [list(r) for r in ws.iter_rows(max_row=HEADER_SEARCH_ROWS, values_only=True)]
        width = max(
            (i for r in head for i, v in enumerate(r, 1) if v is not None and str(v).strip()),
            default=0,
        ) + self._MARGIN
        self._rows = [list(r) for r in ws.iter_rows(max_col=width, values_only=True)]
        self.max_row = len(self._rows)
        self.max_column = width

    def cell(self, row, col):
        values = self._rows[row - 1] if 0 < row <= self.max_row else []
        value = values[col - 1] if 0 < col <= len(values) else None
        return self._Cell(value, row)


def _parse_sheet(ws) -> Offer:
    name_row, label_col = _find_name_row(ws)
    starts = _block_starts(ws, name_row, label_col)

    reference = None
    contractors = []
    last_header_row = name_row
    for index, (start, title) in enumerate(starts):
        end = starts[index + 1][0] - 1 if index + 1 < len(starts) else ws.max_column
        columns = _block_columns(ws, name_row, start, end)
        if columns is None:
            continue
        materials, smr, total, header_row = columns
        last_header_row = max(last_header_row, header_row + 1)
        is_reference = _text(title).startswith(REFERENCE_HEADER)
        if is_reference:
            # Второй блок «Расчетная стоимость» (например, «на 3-ем этапе»)
            # — не эталон, а его вариант: эталон — первый.
            if reference is None:
                reference = Contractor(
                    name="Расчётная стоимость MR Group", materials_col=materials,
                    smr_col=smr, total_col=total, comment_col=None, expected_col=None,
                    header_row=header_row, comment_created=False, top_row=name_row,
                )
            continue

        comment = _labelled_col(ws, name_row, start, end, lambda t: t == COMMENTS_HEADER)
        expected = _labelled_col(ws, name_row, start, end, lambda t: EXPECTED_HEADER in t)
        created = False
        if comment is None and total + 1 <= ws.max_column and _column_is_empty(ws, total + 1):
            comment, created = total + 1, True
        contractors.append(Contractor(
            name=title, materials_col=materials, smr_col=smr, total_col=total,
            comment_col=comment, expected_col=expected,
            header_row=header_row, comment_created=created, top_row=name_row,
        ))
    if not contractors:
        raise OfferError(
            "Не нашлось ни одного блока подрядчика: под названиями в строке "
            "«Наименование контрагента» нет колонок «Стоимость всего» с подписями "
            "«Материалы», «СМР», «Всего»."
        )

    def header_col(predicate):
        return _labelled_col(ws, name_row, 1, label_col + 1, predicate)

    number_col = header_col(lambda t: t == NUMBER_HEADER) or 2
    article_col = header_col(lambda t: t.startswith(ARTICLE_HEADER))
    name_col = header_col(lambda t: t.startswith(WORK_NAME_HEADER))

    lines = []
    section_number = None
    for row in range(last_header_row + 1, ws.max_row + 1):
        raw = ws.cell(row, number_col).value
        # Настоящий номер раздела записан текстом («10.1»): число или
        # результат формулы в этой колонке — случайно попавшее туда
        # количество, а не номер.
        if not isinstance(raw, str):
            continue
        number = raw.strip().rstrip(".")
        level = _level(raw.strip())
        if level is None:
            continue
        if level == LEVEL_SUBSECTION and not (
            section_number and number.startswith(section_number + ".")
        ):
            continue
        name = ""
        for col in (article_col, name_col):
            if col and ws.cell(row, col).value:
                name = " ".join(str(ws.cell(row, col).value).split())
                break
        if LOT_RE.match(name):
            continue
        if level == LEVEL_SECTION:
            section_number = number

        def amounts_of(block):
            return tuple(
                _amount(ws.cell(row, col).value)
                for col in (block.materials_col, block.smr_col, block.total_col)
            )

        lines.append(Line(
            row=row, number=number, name=name, level=level,
            key=estimate_sections.classify(name),
            amounts=[amounts_of(c) for c in contractors],
            reference=amounts_of(reference) if reference else None,
        ))

    has_any_amount = any(
        amount is not None
        for line in lines if line.level == LEVEL_SECTION
        for amounts in line.amounts for amount in amounts
    )
    if not has_any_amount:
        raise OfferError(
            "В файле не нашлось ни одной суммы по разделам. Если таблица собрана "
            "формулами, откройте файл в Excel и сохраните его — тогда суммы "
            "появятся, и файл можно будет загрузить снова."
        )
    has_reference = any(
        line.reference and not _is_empty(line.reference[2])
        for line in lines if line.level == LEVEL_SECTION
    )
    return Offer(
        contractors=contractors, lines=lines, reference=reference,
        has_reference=has_reference,
    )


def project_costs(project):
    """{вид работ: Decimal} — смета объекта плюс подписанное и прогнозируемое
    удорожание по тому же виду работ.

    Подписанное — только если отчёт мерен от сметы (как в «Сводке по
    удорожанию»): иначе его дельта сравнивает файл сам с собой.
    """
    costs = dict(project.estimate or {})
    report = project.signed_report
    if report is not None and report.from_estimate:
        for row in report.rows:
            costs[row.key] = costs.get(row.key, Decimal("0")) + row.delta
    if project.predicted_report is not None:
        for row in project.predicted_report.rows:
            costs[row.key] = costs.get(row.key, Decimal("0")) + row.amount
    return costs


def class_averages(projects, building_class) -> ClassAverages:
    """Средняя ₽/м² по каждому виду работ по объектам ``building_class``.

    Портфельная, как на странице «Сравнить со средним»: сумма стоимостей
    вида работ по объектам, где он есть, делённая на сумму их площадей.
    Объект без сметы или площади не участвует; объект, чью ставку НДС не
    прочитать, выпадает и называется в ``excluded`` — приводить его к 22%
    не от чего.
    """
    adjustments = comparison.Adjustments(vat_rate=OFFER_VAT_RATE)
    totals, areas, counts = {}, {}, {}
    considered = 0
    excluded = []
    for project in projects:
        if (project.passport.get("building_class") or "") != building_class:
            continue
        area = project.passport.get("total_area_sqm") or None
        if not project.estimate or not area:
            continue
        factor = comparison._project_factor_or_exclude(project.passport, adjustments)
        if factor is None:
            excluded.append(project.name)
            continue
        considered += 1
        for key, cost in project_costs(project).items():
            totals[key] = totals.get(key, 0.0) + float(cost) * factor
            areas[key] = areas.get(key, 0.0) + float(area)
            counts[key] = counts.get(key, 0) + 1
    per_sqm = {key: totals[key] / areas[key] for key in totals if areas[key]}
    return ClassAverages(
        per_sqm=per_sqm, counts=counts, considered=considered, excluded=sorted(excluded),
    )


def _is_empty(amount):
    return amount is None or amount < NOMINAL_PRICE_LIMIT


def _unjustified_smr(materials, smr):
    if smr is None or smr <= 0:
        return False
    if not materials or materials <= 0:
        return True
    return smr > materials * SMR_TO_MATERIALS_LIMIT


def _pct(value, base):
    if value is None or not base:
        return None
    return (float(value) / float(base) - 1.0) * 100.0


def _verdict(deviation_pct, section):
    if deviation_pct is None or deviation_pct <= 0:
        return None
    heavy = deviation_pct > HEAVILY_OVERPRICED_PCT
    if section:
        return REMARK_HEAVILY_OVERPRICED if heavy else REMARK_OVERPRICED
    return REMARK_SUB_HEAVILY_OVERPRICED if heavy else REMARK_SUB_OVERPRICED


def _against_reference(verdict, deviation_pct):
    """«Завышена стоимость за раздел, ожидаем снижение на 16,7%» — по
    расчётной стоимости видно не только что, но и насколько."""
    return f"{verdict}, ожидаем снижение на {deviation_pct:.1f}%".replace(".", ",")


def _reference_total(line):
    if line.reference is None or _is_empty(line.reference[2]):
        return None
    return line.reference[2]


def _grouped(lines):
    """Разделы с подразделами под каждым, в порядке файла."""
    groups = []
    for line in lines:
        if line.level == LEVEL_SECTION:
            groups.append((line, []))
        elif groups:
            groups[-1][1].append(line)
    return groups


def _section_view(line, offer, averages, area, with_average):
    """Раздел или подраздел для страницы: сумма каждого подрядчика, ₽/м²,
    отклонения от средней по классу и от расчётной стоимости, оценка."""
    reference = _reference_total(line) if offer.has_reference else None
    avg = averages.per_sqm.get(line.key) if with_average and line.key else None
    expected = avg * area if avg is not None else None
    cells = []
    for _contractor, (_m, _s, total) in zip(offer.contractors, line.amounts):
        if _is_empty(total):
            cells.append(Cell(
                total=total, per_sqm=None, deviation_pct=None, ref_deviation_pct=None,
                remark=REMARK_PRICE_SECTION,
            ))
            continue
        deviation = _pct(total, expected) if expected else None
        ref_deviation = _pct(total, reference) if reference else None
        basis = ref_deviation if offer.has_reference else deviation
        cells.append(Cell(
            total=total, per_sqm=float(total) / area,
            deviation_pct=deviation, ref_deviation_pct=ref_deviation,
            remark=_verdict(basis, section=True),
        ))
    return Section(
        row=line.row, number=line.number, name=line.name, key=line.key,
        label=estimate_sections.CATEGORY_LABELS.get(line.key) if with_average and line.key else None,
        avg_per_sqm=avg,
        count=averages.counts.get(line.key, 0) if avg is not None else 0,
        expected=expected, reference=reference, cells=cells,
    )


def _depth(number):
    return number.count(".")


def analyze(offer, averages, area) -> Analysis:
    """Оценка каждого подрядчика по каждому разделу, который кто-то расценил.

    ``area`` — общая площадь объекта тендера, м².
    """
    area = float(area)
    remarks = []
    sections = []
    groups = []
    for section_line, subsections in _grouped(offer.lines):
        if all(_is_empty(amounts[2]) for amounts in section_line.amounts):
            continue  # раздел вне тендера — его не расценил никто

        section = _section_view(section_line, offer, averages, area, with_average=True)
        sections.append(section)
        for contractor, cell in zip(offer.contractors, section.cells):
            if contractor.comment_col and cell.remark:
                text = cell.remark
                if offer.has_reference and cell.remark != REMARK_PRICE_SECTION:
                    text = _against_reference(cell.remark, cell.ref_deviation_pct)
                remarks.append(Remark(section.row, contractor.comment_col, text))
            if section.expected is not None and contractor.expected_col:
                remarks.append(Remark(
                    section.row, contractor.expected_col, round(section.expected, 2),
                ))

        children = []
        for line in subsections:
            reference = _reference_total(line) if offer.has_reference else None
            for contractor, (materials, smr, total) in zip(offer.contractors, line.amounts):
                if not contractor.comment_col:
                    continue
                if _is_empty(total):
                    remarks.append(Remark(line.row, contractor.comment_col, REMARK_PRICE_SUBSECTION))
                    continue
                parts = []
                if reference:
                    deviation = _pct(total, reference)
                    verdict = _verdict(deviation, section=False)
                    if verdict:
                        parts.append(_against_reference(verdict, deviation))
                if _unjustified_smr(materials, smr):
                    parts.append(REMARK_UNJUSTIFIED_SMR)
                if parts:
                    remarks.append(Remark(line.row, contractor.comment_col, ", ".join(parts)))
            if _depth(line.number) == 1 and not all(_is_empty(a[2]) for a in line.amounts):
                children.append(line)
        groups.append((section, children))

    # По видам работ: разделы, а если сравнение идёт с расчётной стоимостью
    # и раздел в файле один (тендер на одну систему — «ВИС»), то его
    # подразделы: эталон есть и у них, а рейтинг из одной строки «раздел
    # целиком» повторял бы общий.
    if offer.has_reference and len(groups) == 1 and groups[0][1]:
        work_groups = [
            _section_view(line, offer, averages, area, with_average=False)
            for line in groups[0][1]
        ]
    else:
        work_groups = [section for section, _children in groups]

    new_columns = [
        (c.top_row, c.header_row, c.comment_col)
        for c in offer.contractors if c.comment_created
    ]
    return Analysis(
        contractors=[c.name for c in offer.contractors], sections=sections,
        groups=work_groups, remarks=remarks, has_reference=offer.has_reference,
        new_columns=new_columns,
    )


# Одна строка рейтинга: место (None — не расценено), подрядчик, стоимость,
# ₽/м², отклонения от средней по классу и от расчётной стоимости, замечание
# по разделу, сколько разделов не расценено (для общего рейтинга).
RankRow = namedtuple(
    "RankRow",
    "place name total per_sqm deviation_pct ref_deviation_pct remark unpriced",
)
SectionRanking = namedtuple("SectionRanking", "section rows")
# ``reference_total`` — расчётная стоимость MR Group по тем же разделам, или
# None, если её в файле нет.
Ranking = namedtuple("Ranking", "overall expected_total reference_total sections")

# Организационно-правовая форма перед названием — в КП и в паспортах одна и
# та же компания пишется то с ней, то без («АО "ФОДД"» / «ФОДД»).
_LEGAL_FORM_RE = re.compile(
    r"\b(ооо|оао|зао|пао|ао|ип|гк|ук|нао)\b", re.IGNORECASE,
)
_NON_WORD_RE = re.compile(r"[^\w]+")

# Объект подрядчика из базы: название, класс, цена по договору (или None),
# того ли он класса, что выбран для анализа.
HistoryObject = namedtuple("HistoryObject", "name building_class price in_class")
# Сколько наших объектов у подрядчика и на какую сумму — всего и в классе.
ContractorHistory = namedtuple(
    "ContractorHistory", "count total class_count class_total objects",
)


def format_big_money(value):
    """«32,2 млрд ₽» / «850,0 млн ₽» — сумма договоров, которой в колонке
    рейтинга не нужна точность до рубля; «—», если суммы нет."""
    if not value:
        return "—"
    if value >= 1e9:
        return f"{value / 1e9:.1f}".replace(".", ",") + " млрд ₽"
    return f"{value / 1e6:.1f}".replace(".", ",") + " млн ₽"


def normalize_contractor(name):
    """Название подрядчика без формы собственности, кавычек, регистра и
    знаков — чтобы «АО "ФОДД"» из КП и «АО ФОДД» из паспорта были одним."""
    text = str(name or "").lower().replace("ё", "е")
    text = _LEGAL_FORM_RE.sub(" ", text)
    return " ".join(_NON_WORD_RE.sub(" ", text).split())


def contractor_history(names, passports, building_class):
    """{имя подрядчика из КП: ContractorHistory} — объекты, где он у нас
    генподрядчик (поле паспорта ``general_contractor``), с суммой цен по
    договору; отдельно — те же цифры по выбранному классу.

    ``passports`` — список паспортов всех объектов базы.
    """
    by_key = {}
    for passport in passports:
        key = normalize_contractor(passport.get("general_contractor"))
        if key:
            by_key.setdefault(key, []).append(passport)

    result = {}
    for name in names:
        objects = []
        for passport in by_key.get(normalize_contractor(name), []):
            price = passport.get("contract_price_rub")
            objects.append(HistoryObject(
                name=passport.get("project_name") or "—",
                building_class=passport.get("building_class"),
                price=float(price) if isinstance(price, (int, float, Decimal)) else None,
                in_class=passport.get("building_class") == building_class,
            ))
        objects.sort(key=lambda o: (not o.in_class, -(o.price or 0.0)))
        in_class = [o for o in objects if o.in_class]
        result[name] = ContractorHistory(
            count=len(objects),
            total=sum(o.price for o in objects if o.price is not None),
            class_count=len(in_class),
            class_total=sum(o.price for o in in_class if o.price is not None),
            objects=objects,
        )
    return result


def _ranked(rows):
    """Расценённые — от дешёвого к дорогому, с местами; нерасценённые — в
    конце, без места."""
    priced = sorted((r for r in rows if r.total is not None), key=lambda r: r.total)
    unpriced = [r for r in rows if r.total is None]
    return [r._replace(place=i + 1) for i, r in enumerate(priced)] + unpriced


def rank(analysis, area) -> Ranking:
    """Предложения от лучшего к худшему по стоимости — всего и по каждому
    виду работ.

    Общий итог — сумма разделов, которые подрядчик расценил; отклонения от
    средней по классу и от расчётной стоимости считаются только по разделам,
    где есть то, с чем сравнивать, чтобы обе стороны складывались из одного
    и того же. Кто не расценил часть разделов, выглядит дешевле, чем есть, —
    поэтому рядом с ним стоит, сколько разделов не расценено.
    """
    area = float(area)
    expected_total = sum(s.expected for s in analysis.sections if s.expected) or None
    reference_total = sum(
        (s.reference for s in analysis.sections if s.reference), Decimal("0"),
    ) or None
    overall = []
    for index, name in enumerate(analysis.contractors):
        total = Decimal("0")
        vs_average = Decimal("0")
        vs_reference = Decimal("0")
        unpriced = 0
        for section in analysis.sections:
            cell = section.cells[index]
            if cell.remark == REMARK_PRICE_SECTION:
                unpriced += 1
                continue
            total += cell.total
            if section.expected:
                vs_average += cell.total
            if section.reference:
                vs_reference += cell.total
        priced_any = unpriced < len(analysis.sections)
        complete = priced_any and not unpriced
        overall.append(RankRow(
            place=None, name=name,
            total=total if priced_any else None,
            per_sqm=float(total) / area if priced_any else None,
            deviation_pct=_pct(vs_average, expected_total) if complete else None,
            ref_deviation_pct=_pct(vs_reference, reference_total) if complete else None,
            remark=None, unpriced=unpriced,
        ))

    sections = []
    for section in analysis.groups:
        rows = []
        for name, cell in zip(analysis.contractors, section.cells):
            priced = cell.remark != REMARK_PRICE_SECTION
            rows.append(RankRow(
                place=None, name=name,
                total=cell.total if priced else None,
                per_sqm=cell.per_sqm if priced else None,
                deviation_pct=cell.deviation_pct,
                ref_deviation_pct=cell.ref_deviation_pct,
                remark=cell.remark, unpriced=0 if priced else 1,
            ))
        sections.append(SectionRanking(section=section, rows=_ranked(rows)))
    return Ranking(
        overall=_ranked(overall), expected_total=expected_total,
        reference_total=reference_total, sections=sections,
    )


def _frame_created_column(ws, top_row, header_row, col):
    """Колонка «Комментарии», которой в файле не было, — в рамке и заливке
    таблицы: каждая клетка оформлена как соседняя слева на той же строке,
    заголовок занимает обе строки шапки, как подписи у соседей."""
    for row in range(top_row, ws.max_row + 1):
        cell = ws.cell(row, col)
        if not hasattr(cell, "column_letter"):
            continue
        neighbour = ws.cell(row, col - 1)
        if neighbour.has_style:
            cell._style = copy.copy(neighbour._style)
    header = ws.cell(header_row, col)
    header.value = COMMENTS_TITLE
    header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells(start_row=header_row, start_column=col, end_row=header_row + 1, end_column=col)
    ws.column_dimensions[get_column_letter(col)].width = CREATED_COMMENT_WIDTH


def _fit_row(ws, cell, text):
    """Строка таблицы выше, если перенесённый текст замечания в неё не
    помещается — иначе Excel показал бы его обрезанным."""
    lines = math.ceil(len(text) / CHARS_PER_LINE)
    if lines <= 1:
        return
    row = cell.row
    needed = lines * (cell.font.sz or 11) * LINE_SPACING
    current = ws.row_dimensions[row].height
    if current is None or current < needed:
        ws.row_dimensions[row].height = needed


def write_remarks(source_bytes, remarks, new_columns=()) -> bytes:
    """Исходный файл с вписанными ``remarks``, байтами .xlsx.

    Ячейка, где уже что-то есть (замечание, вписанное вручную), не
    перезаписывается. ``new_columns`` — ``(первая строка таблицы, строка
    шапки, колонка)`` колонок «Комментарии», которых в файле не было: они
    встают на место узкого разделителя за блоком подрядчика, получают
    рамку таблицы, заголовок и ширину под текст, а за ними вставляется новый
    разделитель, чтобы следующая компания не прилипала к замечаниям.
    openpyxl не хранит посчитанные значения формул, поэтому книга
    помечается на полный пересчёт при открытии — иначе Excel показал бы в
    ней пустые итоги до первой правки.
    """
    wb = openpyxl.load_workbook(io.BytesIO(source_bytes))
    ws = wb.worksheets[0]
    created = {col for _top, _header, col in new_columns}
    for top_row, header_row, col in new_columns:
        _frame_created_column(ws, top_row, header_row, col)
    for remark in remarks:
        cell = ws.cell(remark.row, remark.col)
        if not hasattr(cell, "column_letter"):
            continue  # не левая верхняя клетка объединённой области
        if cell.value is not None and str(cell.value).strip():
            continue
        cell.value = remark.value
        if remark.col in created:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            _fit_row(ws, cell, str(remark.value))
    xlsx_columns.insert_blank_columns(
        ws, [col + 1 for col in created], width=SPACER_WIDTH,
    )
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
